from __future__ import annotations

import tomllib
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
PROJECT_ROOT = HERE.parents[1]
VERSION = tomllib.loads((PROJECT_ROOT / "source" / "pyproject.toml").read_text(encoding="utf-8"))[
    "project"
]["version"]
ASSETS = HERE / "assets-v5.2"
OUTPUT = HERE / f"KOMI_Report_V{VERSION}_คู่มือติดตั้งและใช้งาน_TH.docx"

FONT = "Prompt"
NAVY = "123262"
BLUE = "0073B7"
ORANGE = "F28C28"
INK = "1F2A3A"
MUTED = "5B6878"
WHITE = "FFFFFF"
BORDER = "D4E0EA"
LIGHT_BLUE = "EDF5FA"
LIGHT_GREEN = "EDF8F2"
LIGHT_ORANGE = "FFF5E8"
LIGHT_RED = "FFF1F0"


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def style_run(run, *, size=10.5, bold=False, color=INK, italic=False) -> None:
    run.font.name = FONT
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    language = properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        properties.append(language)
    for key in ("val", "eastAsia", "bidi"):
        language.set(qn(f"w:{key}"), "th-TH")
    run.font.size = Pt(size)
    run.font.complex_script = True
    run.font.cs_bold = bold
    run.font.cs_italic = italic
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = _rgb(color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, value=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table, color=BORDER, size=6) -> None:
    props = table._tbl.tblPr
    borders = props.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


def text(doc, value="", *, size=10.5, bold=False, color=INK, after=6,
         before=0, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(value)
    style_run(run, size=size, bold=bold, color=color, italic=italic)
    return paragraph


def heading(doc, value: str, *, first=False) -> None:
    if not first:
        doc.add_page_break()
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(value)
    style_run(run, size=19, bold=True, color=NAVY)
    rule = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), BLUE)
    border.append(bottom)
    rule.append(border)


def subheading(doc, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(value)
    style_run(run, size=13.5, bold=True, color=BLUE)


def bullet(doc, value: str, *, level=0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Inches(0.35 + level * 0.22)
    run = paragraph.add_run(value)
    style_run(run, size=10.2)


def callout(doc, title: str, body: str, *, kind="info") -> None:
    palette = {
        "info": (LIGHT_BLUE, NAVY, BLUE),
        "ok": (LIGHT_GREEN, "146B4B", "55A886"),
        "note": (LIGHT_ORANGE, "7A4B00", ORANGE),
        "stop": (LIGHT_RED, "9B1C1C", "D95555"),
    }
    fill, fg, border = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, 150)
    table_borders(table, border, 8)
    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(2)
    style_run(first.add_run(title), size=10.5, bold=True, color=fg)
    second = cell.add_paragraph()
    second.paragraph_format.space_after = Pt(0)
    second.paragraph_format.line_spacing = 1.18
    style_run(second.add_run(body), size=9.7, color=fg)


def figure(doc, filename: str, caption: str, alt: str, *, width=6.35) -> None:
    path = ASSETS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    drawing = run._element.find(qn("w:drawing"))
    if drawing is not None:
        doc_pr = drawing.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("descr", alt)
    text(doc, caption, size=8.3, color=MUTED, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=7)


def checklist(doc, rows: list[str]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(0.65), Inches(5.65))
    table.columns[0].width = widths[0]
    table.columns[1].width = widths[1]
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for index, (cell, value) in enumerate(
        zip(table.rows[0].cells, ("ตรวจ", "รายการ"), strict=True)
    ):
        cell.width = widths[index]
        shade(cell, NAVY)
        cell_margins(cell, 70)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        style_run(paragraph.add_run(value), size=9.3, bold=True, color=WHITE)
    repeat_header(table.rows[0])
    for value in rows:
        cells = table.add_row().cells
        cells[0].width = widths[0]
        cells[1].width = widths[1]
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(cells[0].paragraphs[0].add_run("☐"), size=11, color=BLUE)
        style_run(cells[1].paragraphs[0].add_run(value), size=9.3)
        for cell in cells:
            cell_margins(cell, 65)
    table_borders(table)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(INK)
    fonts = normal._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    for current in doc.sections:
        header = current.header.paragraphs[0]
        header.paragraph_format.space_after = Pt(2)
        style_run(
            header.add_run(f"KOMI Report V{VERSION}  |  คู่มือติดตั้งและใช้งาน"),
            size=8.5,
            bold=True,
            color=MUTED,
        )
        footer = current.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(footer.add_run("หน้า "), size=8.5, color=MUTED)
        run = footer.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, instruction, end))


def cover(doc: Document) -> None:
    logo = ASSETS / "konica-minolta-logo.png"
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(16)
    logo_run = paragraph.add_run()
    logo_run.add_picture(str(logo), width=Inches(2.1))
    for drawing in logo_run._element.findall(".//" + qn("w:drawing")):
        doc_pr = drawing.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("descr", "โลโก้ KOMI Report")
    text(doc, "คู่มือติดตั้งและใช้งาน", size=29, bold=True, color=NAVY, after=4)
    text(doc, f"KOMI Report V{VERSION}", size=20, bold=True, color=BLUE, after=10)
    text(doc, "สำหรับผู้ดูแลสำนักงานและผู้ใช้งานรายงานผ่าน Desktop / LAN Web",
         size=11.5, color=MUTED, after=18)
    callout(
        doc,
        f"คู่มือสำหรับชุดติดตั้ง V{VERSION}",
        "ภาพหน้าจอใช้ข้อมูลตัวอย่างและไม่มีข้อมูลลูกค้า รหัสผ่าน หรือ Token "
        "ชุดติดตั้งที่ไม่มี Authenticode/Developer ID เป็น Unsigned Installation Candidate "
        "ต้องตรวจ SHA-256 ก่อนเปิดไฟล์",
        kind="note",
    )
    figure(
        doc,
        "05-overview.png",
        f"ภาพรวม KOMI Report V{VERSION} หลังตั้งค่าเสร็จ",
        f"หน้าภาพรวม KOMI Report V{VERSION}",
        width=5.55,
    )
    text(doc, f"ฉบับวันที่ {date.today().strftime('%d/%m/%Y')}", size=9, color=MUTED,
         align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)


def start(doc: Document) -> None:
    heading(doc, "เริ่มต้นใช้งาน")
    text(doc, "KOMI Report อ่านข้อมูลเครื่องพิมพ์ Konica Minolta สูงสุด 3 เครื่อง เก็บข้อมูลไว้ใน Database ขององค์กร และเปิดรายงานแบบอ่านอย่างเดียวให้เครื่องอื่นใน LAN ได้")
    subheading(doc, "เตรียมก่อนติดตั้ง")
    bullet(doc, "เครื่อง Windows 10/11 x64 หรือ macOS ที่จะเป็น Primary Desktop")
    bullet(doc, "ชื่อบริษัท และเครื่องพิมพ์ที่อยู่ใน LAN เดียวกัน")
    bullet(doc, "ราคา A4/A3 สีและขาวดำ หรือเลือกตั้งค่าราคาภายหลัง")
    bullet(doc, "Internet สำหรับ Online Registration; Offline Activation ถูกปิดชั่วคราวในรุ่นนี้")
    callout(doc, "Database ควรอยู่ใน Local drive", "ใช้โฟลเดอร์ Documents หรือโฟลเดอร์ข้อมูลภายในเครื่อง หลีกเลี่ยงการวางไฟล์ Database ในโฟลเดอร์ที่ Sync แบบ Real-time ระหว่างโปรแกรมทำงาน", kind="note")
    subheading(doc, "Primary และ Secondary")
    bullet(doc, "Primary เป็นเครื่องแรกที่กำหนดบริษัท เครื่องพิมพ์ ราคา รอบบิล และฐานข้อมูล")
    bullet(doc, "Secondary ใน LAN เดียวกันค้นหา Primary แล้ว Sync การตั้งค่าได้โดยไม่ใช้รหัส Sync หลังผู้ใช้กดอนุญาตบน Primary ภายในช่วงเวลา 5 นาที")
    bullet(doc, "Primary และ Secondary เปิดพร้อมกันได้ แต่ Database ของแต่ละเครื่องเป็นคนละไฟล์")
    bullet(doc, "Web Report ยังคงเปิดดูผ่าน LAN ได้ตามสิทธิ์ที่กำหนด")


def install(doc: Document) -> None:
    heading(doc, "1. ติดตั้งโปรแกรม")
    subheading(doc, "Windows")
    bullet(doc, f"เปิดไฟล์ติดตั้ง KOMI Report V{VERSION} แล้วทำตามหน้าต่าง Setup")
    bullet(doc, "เปิดโปรแกรมจาก Start menu หรือ Desktop shortcut")
    bullet(doc, "ถ้า Windows ขอสิทธิ์ Firewall ให้เลือกเฉพาะ Private network สำหรับ Web Report ใน LAN")
    subheading(doc, "macOS")
    bullet(doc, "เปิดไฟล์ติดตั้ง แล้วลาก KOMI Report ไปยัง Applications")
    bullet(doc, "เปิดจาก Applications และอนุญาต Local Network เมื่อระบบถาม")
    callout(
        doc,
        "ชุดติดตั้งไม่ได้ลงนาม",
        "Windows อาจแสดง Unknown Publisher หรือถูก Smart App Control บล็อก และ macOS อาจถูก Gatekeeper ปฏิเสธ "
        "ตรวจ SHA-256 ให้ตรงกับ CHECKSUMS.sha256 เสมอ หากระบบองค์กรไม่อนุญาตให้หยุดและติดต่อฝ่าย IT "
        "ไม่ควรปิดหรือหลบระบบความปลอดภัย",
        kind="stop",
    )
    callout(doc, "ปิดโปรแกรมจริง", "กดปิดหน้าต่างอาจเป็นการซ่อนลง System Tray/Menu Bar หากต้องการหยุด Service ให้เลือก Exit / ออกจากโปรแกรมจากเมนูไอคอน", kind="info")


def wizard(doc: Document) -> None:
    heading(doc, "2. Setup Wizard — Primary")
    text(doc, "Wizard มี 4 ขั้น และจะลงทะเบียน Online ก่อนเปิดหน้าหลักครั้งแรก")
    figure(doc, "01-setup-company.png", "ขั้นที่ 1 — ชื่อบริษัท วันที่เริ่มใช้ Version และ Database Path", "Setup Wizard ขั้นข้อมูลองค์กร")
    bullet(doc, "กรอกชื่อบริษัท ตรวจวันที่เริ่มใช้งานและ Version")
    bullet(doc, "ใช้ Browse เปลี่ยนตำแหน่ง Database หรือ Import นำฐานข้อมูลเดิม")
    bullet(doc, "อ่านประกาศข้อมูลการลงทะเบียนในกรอบเดียวกัน")
    figure(doc, "02-setup-printers.png", "ขั้นที่ 2 — ค้นหา เลือก และตรวจสอบเครื่องพิมพ์", "Setup Wizard ขั้นค้นหาเครื่องพิมพ์")
    bullet(doc, "ค้นหาอัตโนมัติ หรือใช้กำหนดช่วง IP / เพิ่ม IP ด้วยตนเอง")
    bullet(doc, "เลือกไม่เกิน 3 เครื่อง และตรวจให้ผ่าน SNMP กับ Web")
    bullet(doc, "โปรแกรมลองวิธี Web ตามความสามารถของรุ่น: Modern JSON → Legacy Public XML → Legacy Admin XML เฉพาะเมื่อจำเป็น")
    figure(doc, "03-setup-pricing.png", "ขั้นที่ 3 — ราคา รอบบิล และเปอร์เซ็นต์กระดาษเสีย", "Setup Wizard ขั้นตั้งราคา")
    bullet(doc, "ใช้ราคาเดียวกันทุกเครื่อง ตั้งราคาแยกรายเครื่อง หรือตั้งค่าภายหลัง")
    figure(doc, "04-setup-review.png", "ขั้นที่ 4 — ตรวจสอบทุกส่วนก่อนบันทึก", "Setup Wizard ขั้นตรวจสอบ")
    bullet(doc, "กดบันทึกแล้วรอผล Online Registration")
    bullet(doc, "ถ้าเชื่อม Host ไม่ได้ ให้ตรวจ Internet และการตั้งค่า จากนั้นกดลองใหม่หรือยกเลิกกลับไปตรวจ Wizard")
    callout(
        doc,
        "Secondary Desktop",
        "ถ้าพบ Primary ใน LAN ให้เปิด การตั้งค่าขั้นสูง > รายงานผ่าน LAN บน Primary แล้วกดอนุญาต Secondary Sync 5 นาที "
        "จากนั้น Secondary จึงดึงค่ามาตรวจและบันทึกได้โดยไม่ต้องกรอกรหัส Sync",
        kind="ok",
    )


def daily(doc: Document) -> None:
    heading(doc, "3. อ่านข้อมูลเครื่องพิมพ์")
    figure(doc, "06-devices.png", "หน้าอุปกรณ์ — รุ่น Serial Total Meter หมึก ถาดกระดาษ และ Counter", "หน้าอุปกรณ์ของ KOMI Report")
    bullet(doc, "กด อ่านข้อมูลตอนนี้ เพื่ออ่านทุก Slot โดยไม่รอรอบอัตโนมัติ")
    bullet(doc, "ตรวจชื่อรุ่น Serial Number และ Total Meter ให้ตรงกับเครื่องจริง")
    bullet(doc, "รุ่นที่รองรับจะแสดงระดับหมึก C/Y/M/K และสถานะถาด: Paper Size, Paper Type, Ready/Near Empty/Empty")
    bullet(doc, "การอ่านหลายรุ่นใช้ Targeted SNMP เป็นฐาน และเลือก Modern JSON / Legacy XML ตาม Profile ของเครื่อง โดยไม่ลด TLS ของทุกเครื่องพร้อมกัน")
    bullet(doc, "Counter ต้นทางถูกเก็บสำหรับ Legacy Report และตรวจ Parser Mapping")
    callout(doc, "สถานะไม่ผ่านไม่ได้แปลว่าเครื่องพิมพ์เสียเสมอ", "เปิดสถานะการอ่านเพื่อดูว่าเป็น Network, TLS, Web session, Credential หรือ Parser และแก้เฉพาะสาเหตุ", kind="note")


def reports(doc: Document) -> None:
    heading(doc, "4. รายงานและ Excel")
    figure(doc, "07-reports-overview.png", "Reports มี Overview, Users, Printers, Usage และ Legacy", "หน้ารายงาน Overview")
    bullet(doc, "ทุกหน้ารายงานใช้ช่วงวันที่เดียวกัน")
    bullet(doc, "Overview แยก Print สี/ขาวดำ และ Copy สี/ขาวดำ")
    bullet(doc, "Users แสดงประเภทงานของผู้ใช้แยกตามเครื่อง")
    bullet(doc, "Printers แสดงกราฟทุกเครื่องพร้อมกัน")
    figure(doc, "08-reports-usage.png", "Usage — แยก Print / Copy และดู D, M, Y", "หน้ารายงาน Usage แยก Print และ Copy")
    bullet(doc, "แถวแรกสุดเป็น Meter เริ่มต้น ใช้เป็นฐานอ้างอิง ไม่รวมยอดใช้งานและค่าใช้จ่าย")
    bullet(doc, "คิดรวม % กระดาษเสีย เปิดเป็นค่าเริ่มต้น และปิดได้เพื่อเทียบยอดก่อนส่วนลด")
    figure(doc, "09-reports-legacy.png", "Legacy — แสดง Counter ต้นทางตรง ๆ โดยไม่บังคับ Map เป็น A4/A3", "หน้ารายงาน Legacy")
    bullet(doc, "Legacy แสดง A5, Large Size, FAX และชื่อประเภทอื่นตามที่เครื่องส่งมา")
    bullet(doc, "ดูรวมทุกเครื่อง แยก Slot หรือแยกผู้ใช้ได้")
    subheading(doc, "ส่งออก Excel")
    text(doc, "กด ส่งออก Excel แล้วเลือกช่วงวันที่ ไฟล์มี 6 ชีต:")
    bullet(doc, "การใช้งานรายเครื่อง / มุมมองผู้ใช้ / มุมมองรายวัน")
    bullet(doc, "Legacy รายเครื่อง / Legacy ผู้ใช้ / Legacy รายวัน")
    bullet(doc, "ชีตมาตรฐานแยก Print และ Copy; ชีต Legacy คงชื่อ Counter ต้นทาง")


def settings(doc: Document) -> None:
    heading(doc, "5. Settings และการตั้งค่าขั้นสูง")
    figure(doc, "10-settings.png", "Settings หลักสำหรับเครื่องพิมพ์ ราคา รอบบิล การอ่าน และ Registration", "หน้าตั้งค่าหลัก")
    bullet(doc, "ผู้ใช้ทั่วไปแก้เครื่องพิมพ์ ราคา รอบบิล และตารางอ่านข้อมูลจากหน้าเดียว")
    bullet(doc, "การตั้งค่าขั้นสูงเปิดแสดงข้อมูลภายในหน้า Settings ไม่เด้งไปหน้าใหม่")
    figure(doc, "12-advanced-tools.png", "การตั้งค่าขั้นสูง — Database, Backup, Auto-start, Log และการบำรุงรักษา", "การตั้งค่าขั้นสูง")
    bullet(doc, "ดู/เปลี่ยน Database Path, Backup, Restore และตรวจความสมบูรณ์")
    bullet(doc, "ตั้ง Auto-start, พฤติกรรมปิดหน้าต่าง และระยะเวลาเก็บ Log")
    bullet(doc, "ส่งออก Log สำหรับ Support โดยไม่รวมรหัสผ่านหรือ Token")
    bullet(doc, "ลบ Database / เริ่มระบบใหม่เป็นงานอันตราย ต้องยืนยันและ Backup ก่อน")
    figure(doc, "11-advanced-lan.png", "Web Report ใน LAN อยู่ภายในการตั้งค่าขั้นสูง", "การตั้งค่า Web Report ใน LAN")
    bullet(doc, "Web Report เปิดใน LAN อัตโนมัติ และเลือกเปิดแบบไม่ใช้รหัสได้เฉพาะ Trusted LAN")
    bullet(doc, "ถ้าใช้รหัส ให้เปลี่ยนจาก Desktop; โปรแกรมไม่แสดงรหัสเดิม")
    bullet(doc, "โหมดไม่ใช้รหัสทำให้ผู้ที่อยู่ใน LAN เดียวกันเห็นรายงานและตัวระบุเครื่องพิมพ์ได้ ห้ามเปิด Web Report ออก Public Internet")
    subheading(doc, "Parser Mapping")
    bullet(doc, "เปิดหน้าต่าง ตรวจสอบ Parser Mapping จากการตั้งค่าขั้นสูง")
    bullet(doc, "เลือก Slot แล้วเทียบ Raw Data กับช่อง Database")
    bullet(doc, "แก้ Mapping เฉพาะรายการที่ผิดและบันทึกแยกตาม Serial Number")


def registration(doc: Document) -> None:
    heading(doc, "6. Registration")
    subheading(doc, "Online")
    bullet(doc, "โปรแกรมลงทะเบียนอัตโนมัติหลังจบ Wizard และเก็บ Token ใน OS vault")
    bullet(doc, "ส่ง Status และ Aggregate Total Meter ตามรอบประมาณ 6 ชั่วโมง")
    bullet(doc, "Host อาจแจ้งว่า Desktop ขาดการเชื่อมต่อหลังไม่พบข้อมูลเกิน 3 วัน ไม่ได้สรุปว่าเครื่องพิมพ์เสีย")
    callout(
        doc,
        "Offline Activation ปิดชั่วคราว",
        f"V{VERSION} ไม่แสดงปุ่มส่งออก .komirequest หรือนำเข้า .komilicense ใน Wizard และ Settings "
        "การเปิดใช้งานครั้งแรกจึงต้องลงทะเบียน Online",
        kind="note",
    )
    subheading(doc, "การอัปเดตโปรแกรม")
    bullet(doc, f"V{VERSION} ไม่มีระบบ Update อัตโนมัติ, Forced Update หรือ Remote Command")
    bullet(doc, "เมื่อได้รับรุ่นใหม่ ให้ Exit โปรแกรม, Backup Database, ตรวจ SHA-256 และติดตั้งด้วยไฟล์ชุดใหม่ที่ผู้ส่งมอบอนุมัติ")


def manual_and_support(doc: Document) -> None:
    heading(doc, "7. คู่มือในโปรแกรมและการช่วยเหลือ")
    figure(
        doc,
        "13-manual.png",
        f"เมนูคู่มือในโปรแกรม V{VERSION}",
        "คู่มือการใช้งานภายใน KOMI Report",
    )
    bullet(doc, f"คู่มือในโปรแกรมสรุปการเริ่มต้น วิธีทำงาน รายงาน Settings และระบบ V{VERSION}")
    bullet(doc, "เมื่อต้องขอความช่วยเหลือ ให้ส่ง Version, OS, เวลาที่พบปัญหา, รุ่น/Serial และ Support Log")
    bullet(doc, "ห้ามส่ง Password, Token, Raw credential, เอกสารผู้ใช้ หรือข้อมูลลูกค้าที่ไม่เกี่ยวข้อง")
    subheading(doc, "แก้ปัญหาเร็ว")
    bullet(doc, "ค้นหาไม่พบ: ตรวจ LAN แล้วใช้กำหนดช่วง IP หรือเพิ่ม IP ด้วยตนเอง")
    bullet(doc, "Web ไม่ผ่าน: ตรวจสถานะว่าเป็น certificate, session, login หรือ parser")
    bullet(doc, "ยอดไม่ตรง: ตรวจช่วงวันที่ Meter เริ่มต้น, Hold และ Parser Mapping ก่อนแก้ Database")
    bullet(doc, "Web Report เปิดไม่ได้: ตรวจว่า Desktop เปิดอยู่, Service พร้อม และอยู่ LAN เดียวกัน")


def final_check(doc: Document) -> None:
    heading(doc, "Checklist ก่อนส่งมอบ")
    checklist(doc, [
        "ชื่อบริษัท วันที่เริ่มใช้ Version และ Database Path ถูกต้อง",
        "เลือกเครื่องพิมพ์ไม่เกิน 3 เครื่อง และ Model/Serial ตรงเครื่องจริง",
        "SNMP และ Web ผ่านทุกเครื่องที่เลือก",
        "ราคา รอบบิล และเปอร์เซ็นต์กระดาษเสียถูกต้อง",
        "Online Registration สำเร็จ",
        "อ่าน Total Meter หมึก ถาดกระดาษ และ Counter ได้ตามรุ่นที่รองรับ",
        "รายงานแยก Print/Copy และ Meter เริ่มต้นไม่ถูกนำไปคิดค่าใช้จ่าย",
        "Export Excel เปิดได้ครบ 6 ชีต",
        "Backup Database และ integrity=ok",
        "Web Report เปิดจากเครื่องอื่นใน Trusted LAN ตามสิทธิ์ที่เลือก",
        "ตรวจ SHA-256 ของชุดติดตั้งตรงกับ CHECKSUMS.sha256 และรับทราบว่าไฟล์ไม่ได้ลงนาม",
        "ผู้ใช้รู้วิธี Exit โปรแกรมจริงและเปิด Support Log",
    ])
    callout(doc, "พร้อมใช้งาน", "เมื่อทุกรายการผ่าน ให้เก็บตัวติดตั้ง คู่มือนี้ และข้อมูลติดต่อ Support ไว้ในชุดส่งมอบเดียวกัน", kind="ok")


def build() -> Path:
    required = [
        "konica-minolta-logo.png", "01-setup-company.png", "02-setup-printers.png",
        "03-setup-pricing.png", "04-setup-review.png", "05-overview.png",
        "06-devices.png", "07-reports-overview.png", "08-reports-usage.png",
        "09-reports-legacy.png", "10-settings.png", "11-advanced-lan.png",
        "12-advanced-tools.png", "13-manual.png",
    ]
    missing = [name for name in required if not (ASSETS / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing guide assets: {missing}")
    doc = Document()
    configure(doc)
    props = doc.core_properties
    props.title = f"KOMI Report V{VERSION} - คู่มือติดตั้งและใช้งาน"
    props.subject = "คู่มือลูกค้า Desktop, LAN Web, Reports, Registration และ Backup"
    props.author = "Konica Minolta - KOMI Report Team"
    props.keywords = f"KOMI Report V{VERSION}, Setup, Report, Excel, LAN Web, Online Registration"
    cover(doc)
    start(doc)
    install(doc)
    wizard(doc)
    daily(doc)
    reports(doc)
    settings(doc)
    registration(doc)
    manual_and_support(doc)
    final_check(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    build()
