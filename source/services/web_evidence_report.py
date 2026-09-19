"""Explicit, bounded inspection/export of persisted Web reading intervals."""
from contextlib import closing
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
import json
import sqlite3

from services.web_observations import WebObservationStore
from services.billing_periods import effective_billing_day
from services.web_usage_projection import CounterIdentity, project_counter

MAX_REPORT_EVENTS = 4096
MAX_REPORT_BYTES = 32 * 1024 * 1024

REVIEW_LABELS = {
    "normalization_schema_changed": "รูปแบบการแปลค่ามิเตอร์เปลี่ยน ต้องตรวจฐานใหม่",
    "same_time_counter_conflict": "เวลาอ่านตรงกันแต่ค่ามิเตอร์ต่างกัน",
    "lower_counter_or_clock_conflict": "มิเตอร์ลดลงหรือเวลาของหลักฐานขัดแย้งกัน",
    "counter_baseline_review_required": "ต้องตรวจและยืนยันฐานมิเตอร์",
    "interval_crosses_billing_period": "ช่วงอ่านคร่อมรอบบิล",
    "interval_crosses_price_change": "ช่วงอ่านคร่อมวันเปลี่ยนราคา",
    "selected_counter_missing_on_observed_days": "บางวันมี Web แต่ไม่มีตัวนับที่เลือก",
    "competing_epoch_decisions": "มีการยืนยันฐาน Slot ขัดกัน ต้องตรวจสอบก่อนใช้ยอด",
    "late_parent_evidence": "พบหลักฐานเครื่องเก่าสูงกว่าฐานที่ยืนยันไว้ ต้องตรวจยอดต่อเนื่อง",
    "counter_without_reviewed_baseline": "มีตัวนับที่ยังไม่ได้ยืนยันฐาน",
}


def review_labels(codes):
    return "; ".join(REVIEW_LABELS.get(code, "หลักฐานรอตรวจสอบ") for code in codes)


def load_evidence(database_path, *, slot, serial, start, end, calendar_timezone, cancelled):
    """Old archives are opened only for this explicit inspection request.

    Include no inferred baseline outside the selected interval. The first
    selected observation is the baseline, clearly reported to the reader.
    """
    if type(start) is not date or type(end) is not date or not 0 <= (end - start).days <= 92:
        raise ValueError("เลือกช่วงไม่เกิน 93 วัน")
    first = datetime.combine(start, time.min, calendar_timezone).isoformat()
    last = datetime.combine(end + timedelta(days=1), time.min, calendar_timezone).isoformat()
    store = WebObservationStore(Path(database_path))
    events, size, choices, groups = [], 0, {}, {}
    with closing(sqlite3.connect(store.path)) as conn:
        conn.execute("BEGIN")
        rows = conn.execute("""SELECT event_id,document_json,archive_id FROM web_observations
            WHERE device_slot=? AND device_serial=? AND julianday(observed_at)>=julianday(?)
              AND julianday(observed_at)<julianday(?) ORDER BY sequence LIMIT ?""",
            (slot, serial, first, last, MAX_REPORT_EVENTS + 1))
        for event_id, document, archive_id in rows:
            if cancelled.is_set():
                from concurrent.futures import CancelledError
                raise CancelledError()
            document = store._document(event_id, document, archive_id)
            size += len(document.encode("utf-8"))
            if len(events) >= MAX_REPORT_EVENTS or size > MAX_REPORT_BYTES:
                raise ValueError("ข้อมูลในช่วงนี้มากเกินไป กรุณาเลือกช่วงวันที่สั้นลง")
            event = json.loads(document)
            events.append(event)
            assigned = [item[0] for item in conn.execute("SELECT epoch_id FROM slot_meter_assignments WHERE event_id=?", (event_id,))]
            for epoch_id in assigned or [None]:
                if epoch_id not in groups:
                    epoch = None
                    label = "ข้อมูลเดิมที่ยังไม่ได้ผูกฐานมิเตอร์"
                    if epoch_id is not None:
                        from services.slot_epoch_store import SlotEpochStore
                        epoch = SlotEpochStore._epoch(conn, epoch_id)
                        definition = epoch["document"]
                        action = {"initial": "ฐานเริ่มต้น", "reset": "หลังรีเซ็ต", "replacement": "หลังเปลี่ยนเครื่อง"}[definition["operation"]]
                        at = datetime.fromisoformat(definition["baseline_event"]["payload"]["observed_at"]).astimezone(calendar_timezone)
                        label = f"{action} · {at:%d/%m/%Y %H:%M} · {epoch_id[:8]}"
                    issues = []
                    if epoch is not None:
                        row = conn.execute("SELECT issues_json FROM slot_meter_epoch_state WHERE epoch_id=?", (epoch_id,)).fetchone()
                        issues = json.loads(row[0])
                    if issues:
                        label += " · รอตรวจสอบ"
                    groups[epoch_id] = {"epoch": epoch, "label": label, "events": [], "issues": issues}
                groups[epoch_id]["events"].append(event)
            for c in event["payload"]["reading"]["counters"]:
                if not c["canonical_field"].startswith("counter."):
                    continue
                key = CounterIdentity(c["canonical_field"], c["subject_type"], c["subject_id"])
                subject = f" · {c['subject_name'] or c['subject_id']} ({c['subject_id']})" if c["subject_type"] else ""
                choices[key] = f"{c['label']}{subject}"
        prices = [date.fromisoformat(row[0]) for row in conn.execute("SELECT DISTINCT effective_date FROM price_history WHERE device_slot=?", (slot,))]
        row = conn.execute("SELECT value FROM settings WHERE key='billing_day'").fetchone()
        billing_day = effective_billing_day(row[0] if row else None)
        current = conn.execute("SELECT epoch_id FROM slot_meter_heads WHERE device_slot=?", (slot,)).fetchone()
    return {"events": events, "choices": sorted(choices.items(), key=lambda item: item[1]),
            "price_change_dates": prices, "billing_day": billing_day, "epochs": groups,
            "current_epoch": current[0] if current and current[0] in groups else None}


def inspect_counter(evidence, counter, *, slot, serial, start, end, calendar_timezone, epoch_id=None):
    # Inspection is not approval of a reset epoch or a billable posting.
    groups = evidence.get("epochs", {None: {"events": evidence["events"], "epoch": None}})
    if epoch_id is None and len(groups) == 1:
        epoch_id = next(iter(groups))
    # The UI passes 'legacy' explicitly so choosing it cannot silently select
    # a different group after an asynchronous load.
    if epoch_id == "legacy":
        epoch_id = None
    elif epoch_id is None and len(groups) > 1:
        raise ValueError("เลือกช่วงฐานมิเตอร์ก่อน ไม่รวมข้อมูลก่อนและหลังรีเซ็ต")
    if epoch_id not in groups:
        raise ValueError("ไม่พบข้อมูลในช่วงฐานมิเตอร์ที่เลือก")
    group = groups[epoch_id]
    result = project_counter(group["events"], slot=slot, serial=serial,
        epoch_id=epoch_id or f"inspection:{serial}:{start}:{end}", counter=counter,
        reviewed_epoch=group["epoch"],
        calendar_timezone=calendar_timezone, billing_day=evidence["billing_day"],
        price_change_dates=evidence["price_change_dates"])
    issues = set(group.get("issues", ()))
    if issues:
        result = replace(result, issues=tuple(sorted(set(result.issues) | issues)),
            intervals=tuple(replace(item, holds=tuple(sorted(set(item.holds) | issues))) for item in result.intervals))
    return result


def export_evidence_xlsx(path, projection, *, printer, serial, counter_label, cancelled=None, progress=None):
    from concurrent.futures import CancelledError
    from i18n import tr
    def checkpoint(percent, text):
        if cancelled is not None and cancelled.is_set():
            raise CancelledError()
        if progress:
            progress(percent, text)
    checkpoint(5, tr("สร้างสมุดงาน…", "Creating workbook…"))
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from services.report_excel_export import _neutralize_untrusted_cell_string
    workbook = Workbook()
    summary = workbook.active
    summary.title = "ขอบเขตข้อมูล"
    for row in [
        ["เครื่องพิมพ์", printer], ["Serial", serial], ["มิเตอร์", counter_label],
        ["รุ่นหลักฐาน", projection.revision],
        ["ช่วงฐานมิเตอร์", projection.epoch_id],
        ["ขอบเขต", "เริ่มนับจากจุด Web แรกในช่วงที่เลือก; รายงานหลักฐานนี้ไม่ใช่ยอดเรียกเก็บหรือการยืนยันฐานหลัง reset"],
        ["กราฟประมาณการ", "เฉลี่ยเฉพาะวันที่ไม่มี Web ทั้งวัน ปัดเฉพาะการแสดง ยอดบันทึกวันก่อนกลับมาอ่านได้ไม่ถูกบวกซ้ำ"],
        ["ข้อควรตรวจ", review_labels(projection.issues)],
    ]:
        summary.append([_neutralize_untrusted_cell_string(str(value))[0] for value in row])
    points = workbook.create_sheet("อ่านจริง")
    points.append(["เวลาอ่านจริง", "มิเตอร์สะสม", "Event IDs"])
    for index, point in enumerate(projection.points):
        if index % 100 == 0:
            checkpoint(10 + int(30 * index / max(1, len(projection.points))), tr("เขียนค่ามิเตอร์…", "Writing meter readings…"))
        points.append([point.at.isoformat(), float(point.value), ", ".join(point.evidence_ids)])
    intervals = workbook.create_sheet("ยอดตามช่วง")
    intervals.append(["เริ่ม", "ถึง", "ผลต่าง", "วันลงยอด", "วันที่ไม่มี Web", "รอตรวจ", "หมายเหตุ"])
    for index, item in enumerate(projection.intervals):
        if index % 100 == 0:
            checkpoint(40 + int(30 * index / max(1, len(projection.intervals))), tr("เขียนช่วงหลักฐาน…", "Writing evidence intervals…"))
        intervals.append([item.start.at.isoformat(), item.end.at.isoformat(), float(item.amount), item.allocation_day.isoformat(),
            ", ".join(day.isoformat() for day in item.empty_days), review_labels(item.holds), item.remark])
    graph = workbook.create_sheet("กราฟรายวัน")
    graph.append(["วันที่", "ยอดช่วงมีข้อมูล", "ยอดประมาณการ", "ยอดแสดง", "เป็นค่าประมาณ", "หมายเหตุ"])
    checkpoint(75, tr("เขียนกราฟรายวัน…", "Writing daily graph data…"))
    for item in projection.daily_graph():
        graph.append([item.day.isoformat(), float(item.observed_interval_usage), float(item.estimated_usage),
                      float(item.total), "ประมาณการ" if item.estimated else "", "\n".join(item.remarks)])
        for column in (2, 3, 4):
            graph.cell(graph.max_row, column).number_format = "#,##0.##"
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in "ABCDEFG":
            sheet.column_dimensions[column].width = 28 if column != "G" else 70
    checkpoint(85, tr("บีบอัดไฟล์ Excel…", "Compressing Excel file…"))
    try:
        workbook.save(path)
    finally:
        workbook.close()
    checkpoint(90, tr("ตรวจไฟล์เสร็จแล้ว", "File writing complete"))
