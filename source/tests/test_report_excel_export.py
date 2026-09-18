from dataclasses import replace
from datetime import UTC, date, datetime

from openpyxl import Workbook, load_workbook

from services.device_config import DeviceConfig
from services.monthly_history import MonthlyUsage
from services.raw_counter_history import RawCounterDailyRecord
from services.report_excel_export import (
    OPERATION_DEVICE_SUBHEADERS,
    CloudReportsExportData,
    DailyCategoryCost,
    _set_cell_value,
    export_cloud_reports_xlsx,
)
from services.subject_history import SubjectPeriodUsage


def _usage(day: date, *, slot: int, identity: str, name: str, total: float) -> MonthlyUsage:
    return MonthlyUsage(
        device_slot=slot,
        period_start=day,
        period_end=day,
        total=total,
        a3_color=10 + slot,
        a3_bw=5 + slot,
        a4_color=50 + slot,
        a4_bw=55 + slot,
        scan=12 + slot,
        gross_cost=30 + slot,
        net_cost=29.4 + slot,
        day_count=1,
        incomplete_days=0,
        device_identity=identity,
        device_name_snapshot=name,
        device_serial_snapshot=identity.removeprefix("serial:"),
        print_color=48 + slot,
        print_bw=30 + slot,
        copy_color=12 + slot,
        copy_bw=30 + slot,
        print_a4_color=40 + slot,
        print_a4_bw=25 + slot,
        print_a3_color=8,
        print_a3_bw=5,
        copy_a4_color=10,
        copy_a4_bw=30,
        copy_a3_color=2 + slot,
        copy_a3_bw=0,
        print_available=True,
        copy_available=True,
    )


def _subject(
    day: date,
    *,
    slot: int,
    identity: str,
    name: str,
    usage: float,
) -> SubjectPeriodUsage:
    return SubjectPeriodUsage(
        device_slot=slot,
        period_start=day,
        period_end=day,
        subject_type="user",
        subject_id="u-1",
        subject_name="alice",
        latest_total=500,
        period_usage=usage,
        reset_count=0,
        missing_baseline_days=0,
        incomplete_days=0,
        sample_count=2,
        latest_at="2026-08-06T09:00:00+07:00",
        warnings=(),
        device_identity=identity,
        device_name_snapshot=name,
        device_serial_snapshot=identity.removeprefix("serial:"),
    )


def _raw(
    day: date,
    *,
    slot: int,
    field: str,
    label: str,
    delta: float,
    subject: bool = False,
    serial: str | None = None,
    name: str | None = None,
) -> RawCounterDailyRecord:
    return RawCounterDailyRecord(
        device_slot=slot,
        record_date=day,
        canonical_field=field,
        label=label,
        source_path="/wcd/system_counter.xml",
        subject_type="user" if subject else "",
        subject_id="u-1" if subject else "",
        subject_name="alice" if subject else "",
        opening_value=100,
        closing_value=100 + delta,
        delta_value=delta,
        reset_count=0,
        sample_count=2,
        first_success="2026-08-06T08:00:00+07:00",
        last_success="2026-08-06T09:00:00+07:00",
        source_name="legacy_xml",
        device_name_snapshot=name or f"MFP {slot + 1}",
        device_serial_snapshot=serial or ("SER-A", "SER-B", "SER-C")[slot],
        incomplete=False,
    )


def test_cloud_reports_export_writes_six_wide_safe_views(tmp_path):
    day = date(2026, 8, 6)
    first = _usage(day, slot=0, identity="serial:SER-A", name="MFP 1", total=120)
    second = _usage(day, slot=1, identity="serial:SER-B", name="MFP 2", total=80)
    output = export_cloud_reports_xlsx(
        tmp_path / "reports",
        CloudReportsExportData(
            selected_records=(first, second),
            selected_daily_records=(first, second),
            usage_daily_records=(first, second),
            monthly_records=(first, second),
            subjects=(
                _subject(day, slot=0, identity="serial:SER-A", name="MFP 1", usage=75),
                _subject(day, slot=1, identity="serial:SER-B", name="MFP 2", usage=25),
            ),
            devices=(
                DeviceConfig(
                    0,
                    "MFP 1",
                    "192.0.2.10",
                    username="administrator",
                    password="do-not-export-this",
                    serial="SER-A",
                ),
                DeviceConfig(
                    1,
                    "MFP 2",
                    "192.0.2.11",
                    username="administrator",
                    password="also-private",
                    serial="SER-B",
                ),
            ),
            period=("2026-08-01", "2026-08-31"),
            last_sync_at="2026-08-06T08:50:00+07:00",
            generated_at=datetime(2026, 8, 6, 8, 55, tzinfo=UTC),
            category_costs=(
                DailyCategoryCost(0, "serial:SER-A", day, 10, 5, 8, 6.4),
                DailyCategoryCost(1, "serial:SER-B", day, 11, 5, 8, 6.4),
            ),
            raw_counter_records=(
                _raw(day, slot=0, field="paper.a5", label="A5", delta=7),
                _raw(day, slot=1, field="fax.total", label="FAX", delta=2),
                _raw(day, slot=0, field="paper.a5", label="A5", delta=5, subject=True),
            ),
        ),
    )

    assert output.name == "reports.xlsx"
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == [
        "การใช้งานรายเครื่อง",
        "มุมมองผู้ใช้",
        "มุมมองรายวัน",
        "Legacy-รายเครื่อง",
        "Legacy-ผู้ใช้",
        "Legacy-รายวัน",
    ]

    machines = workbook["การใช้งานรายเครื่อง"]
    assert machines["A5"].value == "MFP 1"
    assert machines["B3"].value == "Print A4 สี"
    assert machines["B5"].value == 40
    assert machines["L3"].value == "Copy A4 สี"
    assert machines["L5"].value == 10
    assert machines["W3"].value == "Print"
    assert machines["W5"].value == 78
    assert machines["X3"].value == "Copy"
    assert machines["X5"].value == 42
    assert machines["Y5"].value == 120
    assert machines["Z5"].value == 29.4
    assert machines.freeze_panes == "B5"

    users = workbook["มุมมองผู้ใช้"]
    device_block = len(OPERATION_DEVICE_SUBHEADERS)
    assert users["A5"].value == "alice"
    assert users.cell(4, 4).value == "Print A4 สี (หน้า)"
    assert users.cell(4, 4 + 10).value == "Copy A4 สี (หน้า)"
    assert users.cell(5, 4 + device_block - 2).value == 75
    assert users.cell(5, 4 + device_block + device_block - 2).value == 25
    assert users.cell(5, 3 + 2 * device_block + 1).value == 100
    assert users["D5"].value is None
    assert users.freeze_panes == "D5"

    daily = workbook["มุมมองรายวัน"]
    assert daily["A5"].value.date() == day
    assert daily.max_row == 5
    assert daily["B5"].value == 40
    assert daily.cell(5, 2 + 10).value == 10
    assert daily.cell(5, 2 + device_block).value == 41
    assert daily.cell(5, 1 + 2 * device_block + 1).value == 200
    assert daily.freeze_panes == "B5"

    legacy_machine = workbook["Legacy-รายเครื่อง"]
    assert legacy_machine["A4"].value == "กลุ่มงาน"
    assert legacy_machine["B4"].value == "ประเภท Counter (ชื่อจากเครื่อง)"
    assert legacy_machine["B5"].value == "A5"
    assert legacy_machine["C5"].value == 7
    assert legacy_machine["D5"].value == "—"
    assert legacy_machine["C6"].value == "—"
    assert legacy_machine["D6"].value == 2
    assert legacy_machine.freeze_panes == "C5"

    legacy_users = workbook["Legacy-ผู้ใช้"]
    assert legacy_users["A5"].value == "alice (user)"
    assert legacy_users["B5"].value == 5
    assert legacy_users.freeze_panes == "B5"

    legacy_daily = workbook["Legacy-รายวัน"]
    assert legacy_daily["A5"].value.date() == day
    assert legacy_daily["B5"].value == 7
    assert legacy_daily["C5"].value == 2
    assert legacy_daily.freeze_panes == "B5"

    all_text = "\n".join(
        str(cell.value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "do-not-export-this" not in all_text
    assert "also-private" not in all_text
    assert "192.0.2.10" not in all_text
    assert "SER-A" not in all_text


def test_excel_can_export_cost_before_waste_adjustment(tmp_path):
    day = date(2026, 8, 6)
    record = _usage(day, slot=0, identity="serial:SER-A", name="MFP 1", total=120)
    output = export_cloud_reports_xlsx(
        tmp_path / "gross-cost.xlsx",
        CloudReportsExportData(
            selected_records=(record,),
            selected_daily_records=(record,),
            usage_daily_records=(record,),
            monthly_records=(record,),
            subjects=(),
            devices=(DeviceConfig(0, "MFP 1", "192.0.2.10", serial="SER-A"),),
            period=(day.isoformat(), day.isoformat()),
            last_sync_at=None,
            generated_at=datetime(2026, 8, 6, 8, 55, tzinfo=UTC),
            include_waste_adjustment=False,
        ),
    )

    workbook = load_workbook(output, data_only=False)
    assert workbook["การใช้งานรายเครื่อง"]["Z5"].value == 30
    assert "ก่อนหักกระดาษเสีย" in workbook["การใช้งานรายเครื่อง"]["A2"].value


def test_daily_view_has_one_row_per_date_even_with_multiple_printers(tmp_path):
    first_day = date(2026, 8, 6)
    second_day = date(2026, 8, 7)
    records = (
        _usage(first_day, slot=0, identity="serial:A", name="MFP 1", total=10),
        _usage(first_day, slot=1, identity="serial:B", name="MFP 2", total=20),
        _usage(second_day, slot=0, identity="serial:A", name="MFP 1", total=30),
    )
    output = export_cloud_reports_xlsx(
        tmp_path / "daily.xlsx",
        CloudReportsExportData(
            selected_records=records,
            selected_daily_records=records,
            usage_daily_records=records,
            monthly_records=(),
            subjects=(),
            devices=(),
            period=(first_day.isoformat(), second_day.isoformat()),
            last_sync_at=None,
            generated_at=datetime.now(UTC),
        ),
    )
    daily = load_workbook(output)["มุมมองรายวัน"]
    assert daily.max_row == 6
    assert [daily.cell(row, 1).value.date() for row in (5, 6)] == [first_day, second_day]
    second_device_start = 2 + len(OPERATION_DEVICE_SUBHEADERS)
    assert all(
        daily.cell(6, column).value is None
        for column in range(
            second_device_start,
            second_device_start + len(OPERATION_DEVICE_SUBHEADERS),
        )
    )


def test_legacy_sheets_separate_replacement_printers_in_same_slot(tmp_path):
    day = date(2026, 8, 6)
    records = (
        _raw(
            day,
            slot=0,
            field="paper.a5",
            label="A5",
            delta=5,
            serial="OLD-0001",
            name="Old Printer",
        ),
        _raw(
            day,
            slot=0,
            field="paper.a5",
            label="A5",
            delta=7,
            serial="NEW-0002",
            name="New Printer",
        ),
    )
    output = export_cloud_reports_xlsx(
        tmp_path / "legacy-identity.xlsx",
        CloudReportsExportData(
            selected_records=(),
            selected_daily_records=(),
            usage_daily_records=(),
            monthly_records=(),
            subjects=(),
            devices=(),
            period=(day.isoformat(), day.isoformat()),
            last_sync_at=None,
            generated_at=datetime.now(UTC),
            raw_counter_records=records,
        ),
    )

    workbook = load_workbook(output)
    machine = workbook["Legacy-รายเครื่อง"]
    assert machine["C4"].value != machine["D4"].value
    assert "Old Printer" in machine["D4"].value or "Old Printer" in machine["C4"].value
    assert "New Printer" in machine["D4"].value or "New Printer" in machine["C4"].value
    assert {machine["C5"].value, machine["D5"].value} == {5, 7}

    daily = workbook["Legacy-รายวัน"]
    assert daily["B4"].value != daily["C4"].value
    assert {daily["B5"].value, daily["C5"].value} == {5, 7}

    all_text = "\n".join(
        str(cell.value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "OLD-0001" not in all_text
    assert "NEW-0002" not in all_text


def test_excel_export_neutralizes_untrusted_formula_strings(tmp_path):
    day = date(2026, 8, 6)
    device_name = "\x00 \t=HYPERLINK(\"https://example.invalid\",\"open\")"
    user_name = "\t@SUM(1,1)"
    raw_label = "\ufeff+CMD|' /C calc'!A0"
    record = replace(
        _usage(day, slot=0, identity="serial:SER-A", name=device_name, total=120),
        device_name_snapshot=device_name,
    )
    subject = replace(
        _subject(
            day,
            slot=0,
            identity="serial:SER-A",
            name=device_name,
            usage=25,
        ),
        subject_type="+user",
        subject_id=" -1+1",
        subject_name=user_name,
    )
    device_raw = replace(
        _raw(
            day,
            slot=0,
            field="custom.counter",
            label=raw_label,
            delta=-7,
            name=device_name,
        ),
        source_path="\r=external.path",
    )
    subject_raw = replace(
        _raw(
            day,
            slot=0,
            field="custom.subject",
            label="-External label",
            delta=4,
            subject=True,
            name=device_name,
        ),
        subject_type=subject.subject_type,
        subject_id=subject.subject_id,
        subject_name=subject.subject_name,
        source_path="\n@external.path",
    )

    output = export_cloud_reports_xlsx(
        tmp_path / "formula-safe.xlsx",
        CloudReportsExportData(
            selected_records=(record,),
            selected_daily_records=(record,),
            usage_daily_records=(record,),
            monthly_records=(record,),
            subjects=(subject,),
            devices=(
                DeviceConfig(
                    0,
                    device_name,
                    "192.0.2.10",
                    serial="SER-A",
                ),
            ),
            period=(day.isoformat(), day.isoformat()),
            last_sync_at=None,
            generated_at=datetime.now(UTC),
            raw_counter_records=(device_raw, subject_raw),
        ),
    )

    workbook = load_workbook(output, data_only=False)
    dangerous_fragments = (
        "HYPERLINK",
        "@SUM(1,1)",
        "CMD|' /C calc'!A0",
        "-External label",
        "+user",
        "-1+1",
    )
    matched = {
        fragment: []
        for fragment in dangerous_fragments
    }
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                for fragment in dangerous_fragments:
                    if fragment in value:
                        matched[fragment].append(cell)

    assert all(matched.values())
    for cells in matched.values():
        for cell in cells:
            significant = str(cell.value).lstrip("' \t\r\n\ufeff")
            if significant[0] in "=+-@":
                assert str(cell.value).startswith("'")
                assert cell.data_type == "s"
                assert cell.quotePrefix is True

    legacy_machine = workbook["Legacy-รายเครื่อง"]
    negative_numeric_cells = [
        cell
        for row in legacy_machine.iter_rows()
        for cell in row
        if cell.value == -7
    ]
    assert negative_numeric_cells
    assert all(cell.data_type == "n" for cell in negative_numeric_cells)
    assert isinstance(workbook["มุมมองรายวัน"]["A5"].value, datetime)


def test_user_cost_uses_each_days_evidence_and_keeps_missing_cost_unavailable(tmp_path):
    from excel_export_fixture import report_for
    report = report_for(users=1, days=2, printers=1)
    raw, costs = [], []
    for day, printer_counts, user_counts, cost in (
        (date(2026, 7, 1), (10, 10), (3, 1), 40),
        (date(2026, 7, 2), (30, 10), (2, 2), 120),
    ):
        for operation, device_pages, user_pages in zip(('print', 'copy'), printer_counts, user_counts):
            for subject, pages in ((False, device_pages), (True, user_pages)):
                raw.append(replace(_raw(day, slot=0, field=f'counter.{operation}.full_color.a4',
                    label=f'{operation} A4 color', delta=pages, subject=subject, serial='SYNTH-0'),
                    subject_id='0' if subject else ''))
        costs.append(DailyCategoryCost(0, 'serial:SYNTH-0', day, cost, 0, 0, 0))
    report = replace(report, raw_counter_records=tuple(raw), category_costs=tuple(costs),
                     subjects=(replace(report.subjects[0], period_usage=8),))
    for missing in (False, True):
        payload = replace(report, category_costs=tuple(costs[:1])) if missing else report
        book = load_workbook(export_cloud_reports_xlsx(tmp_path/f'cost-{missing}.xlsx', payload), data_only=True)
        try:
            users = book['มุมมองผู้ใช้']
            assert users['D5'].value == 5 and users['N5'].value == 3
            # Rates differ each day: Print = 3*2 + 2*3; Copy = 1*2 + 2*3.
            assert users['E5'].value == (None if missing else 12)
            assert users['O5'].value == (None if missing else 8)
            assert users['AB5'].value == (None if missing else 20)
            assert users['AA5'].value == 8
        finally:
            book.close()


def test_excel_cell_writer_preserves_program_authored_formula():
    workbook = Workbook()
    cell = _set_cell_value(
        workbook.active,
        1,
        1,
        "=SUM(1,2)",
        trusted_formula=True,
    )

    assert cell.value == "=SUM(1,2)"
    assert cell.data_type == "f"
