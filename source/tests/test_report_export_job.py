from concurrent.futures import CancelledError, ThreadPoolExecutor
from datetime import date
from threading import Event, get_ident
from time import monotonic

import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import QTimer
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication, QPushButton, QWidget

import database
from excel_export_fixture import report_for
from services.report_excel_export import export_cloud_reports_xlsx
from ui.report_export import ReportExportController


def wait_until(predicate, seconds=15):
    deadline = monotonic() + seconds
    while not predicate() and monotonic() < deadline:
        QTest.qWait(10)
    assert predicate()


def test_90_users_export_keeps_qt_responsive_and_verifies_workbook(tmp_path):
    # Timing measurements need their own QApplication/process: the functional
    # suite deliberately creates and tears down hundreds of unrelated Qt
    # windows. Keep the timer/click thresholds and workbook assertions intact.
    import os
    import subprocess
    import sys
    if os.environ.get("KOMI_TEST_FRESH_QT_EXPORT") != "1":
        result = subprocess.run(
            [sys.executable, "-m", "pytest", __file__ + "::test_90_users_export_keeps_qt_responsive_and_verifies_workbook",
             "-q", "--tb=short", "--basetemp", str(tmp_path / "fresh-qt")],
            env={**os.environ, "KOMI_TEST_FRESH_QT_EXPORT": "1"},
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
            encoding="utf-8", errors="replace", timeout=120,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        assert result.returncode == 0, result.stdout
        return
    app = QApplication.instance() or QApplication([])
    from config import build_qss
    from ui_consistency import install_button_system
    install_button_system(app)
    app.setStyleSheet(build_qss())
    parent = QWidget()
    button = QPushButton('Navigate', parent)
    parent.show()
    controller = ReportExportController(parent)
    clicks, ticks, successes, errors = [], [], [], []
    button.clicked.connect(lambda: clicks.append(True))
    controller.succeeded.connect(successes.append)
    controller.failed.connect(errors.append)
    timer = QTimer(parent)
    timer.setInterval(10)
    timer.timeout.connect(lambda: ticks.append(monotonic()))
    timer.timeout.connect(button.click)
    timer.start()
    report = report_for()
    updates = []
    def export(*, cancelled, progress):
        def notify(percent, message):
            updates.append((percent, message))
            progress(percent, message)
        return export_cloud_reports_xlsx(tmp_path/'users.xlsx', report, cancelled=cancelled, progress=notify)
    try:
        controller.start(export)
        assert not controller.start(lambda **kwargs: None)
        wait_until(lambda: not controller.busy, seconds=45)
        assert successes and not errors and len(ticks) > 10 and len(clicks) == len(ticks)
        # Only a coarse ceiling for a frozen GUI; performance timings are in the benchmark.
        assert max(b - a for a, b in zip(ticks, ticks[1:])) < 2
        assert [value for value, _ in updates] == sorted(value for value, _ in updates)
        assert any('45/90 คน' in message for _, message in updates)
        assert any('90/90 คน' in message for _, message in updates)
        book = load_workbook(successes[0], read_only=True, data_only=True)
        rows = list(book['มุมมองผู้ใช้'].iter_rows(min_row=5, values_only=True))
        assert len(rows) == 90
        assert all(row[-2] == 720 and row[27] == 480 for row in rows)
        assert book['Legacy-ผู้ใช้'].max_row == 94
        assert book['มุมมองรายวัน'].max_row == 34
        book.close()
    finally:
        timer.stop()
        controller.shutdown()
        parent.close()
        app.processEvents()


@pytest.mark.parametrize('action', ['cancel', 'fail', 'locked'])
def test_incomplete_export_preserves_destination_and_removes_temporary_files(tmp_path, monkeypatch, action):
    import services.report_excel_export as excel
    destination = tmp_path/'existing.xlsx'
    destination.write_bytes(b'original workbook')
    cancelled = Event()
    if action == 'locked':
        def locked(*args):
            raise PermissionError('destination is open in Excel')
        monkeypatch.setattr(excel.os, 'replace', locked)
    def stage(percent, message):
        if percent >= 25 and action != 'locked':
            if action == 'fail':
                raise OSError('synthetic failure')
            cancelled.set()
    with pytest.raises(CancelledError if action == 'cancel' else OSError):
        excel.export_cloud_reports_xlsx(destination, report_for(users=1, days=1),
                                      cancelled=cancelled, progress=stage)
    assert destination.read_bytes() == b'original workbook'
    assert list(tmp_path.glob('.komi-export-*')) == []


def test_controller_cancel_then_retry_and_failure_restore_buttons():
    app = QApplication.instance() or QApplication([])
    parent = QWidget()
    controller = ReportExportController(parent)
    started, release = Event(), Event()
    states, errors, successes = [], [], []
    controller.busy_changed.connect(states.append)
    controller.failed.connect(errors.append)
    controller.succeeded.connect(successes.append)
    main = get_ident()
    def blocked(*, cancelled, progress):
        assert get_ident() != main
        progress(47, 'กำลังสร้างรายงานผู้ใช้ 2/3 คน…')
        started.set()
        release.wait(5)
        if cancelled.is_set():
            raise CancelledError()
    try:
        controller.start(blocked)
        wait_until(started.is_set)
        dialog = controller.dialog
        wait_until(lambda: dialog.progress_bar.value() == 47)
        assert dialog.isVisible() and '2/3' in dialog.phase_label.text()
        initial_elapsed = dialog.elapsed_label.text()
        wait_until(lambda: dialog.elapsed_label.text() != initial_elapsed)
        # Cancel and window-close keep the popup visible until the worker stops.
        dialog.cancel_button.click()
        assert controller.cancelled.is_set() and controller.busy and dialog.isVisible()
        assert not dialog.cancel_button.isEnabled()
        dialog.close()
        assert dialog.isVisible() and 'กำลังยกเลิก' in dialog.phase_label.text()
        release.set()
        wait_until(lambda: not controller.busy)
        assert states == [True, False] and not errors and not successes
        def fail(**kwargs):
            raise OSError('permission denied')
        controller.start(fail)
        wait_until(lambda: not controller.busy)
        assert len(errors) == 1 and states[-1] is False
        controller.start(lambda **kwargs: 'retry succeeded')
        wait_until(lambda: not controller.busy)
        assert successes == ['retry succeeded']
    finally:
        release.set()
        controller.shutdown()
        parent.close()
        app.processEvents()


def test_background_read_database_is_pinned_readonly_and_does_not_leak_settings(tmp_path):
    first = database.get_db_path()
    database.set_cfg('billing_day', '7')
    started, release = Event(), Event()
    second = tmp_path/'second.db'
    def read():
        with database.report_read_database(first):
            started.set()
            release.wait(3)
            assert database.get_db_path() == first.resolve()
            assert database.get_cfg('billing_day') == '7'
            with pytest.raises(Exception, match='readonly'):
                database.set_cfg('billing_day', '1')
            return database.get_cfg('billing_day')
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(read)
        assert started.wait(2)
        try:
            database.set_db_path(second)
            database.init_db()
            database.set_cfg('billing_day', '25')
            release.set()
            assert future.result(3) == '7'
            assert database.get_db_path() == second and database.get_cfg('billing_day') == '25'
        finally:
            release.set()
            database.set_db_path(first)


@pytest.mark.parametrize('archived', [False, True])
def test_export_job_reads_selected_user_and_raw_dates_from_database(tmp_path, archived):
    from services.report_export_job import export_report_selection, EmptyReportError
    from services.device_config import DeviceConfig, DeviceConfigRepository
    from services.pricing import PricingRepository, PriceList
    from services.billing_ledger import post_daily_record
    from test_billing_ledger import _daily
    device = DeviceConfig(0, 'Fixture', '192.0.2.10', serial='TEST')
    DeviceConfigRepository().replace_all([device])
    PricingRepository().set_prices(0, date(2026, 7, 1), PriceList(a4_color=1))
    for day, pages in (('2026-07-01', 999), ('2026-07-02', 3), ('2026-07-03', 888)):
        with database._open() as conn:
            ledger_source = _daily(conn, {'reconciliation': {'billing_safe': True}}, record_date=day)
            conn.execute('''INSERT INTO daily_subject_records(device_slot,record_date,subject_type,subject_id,subject_name,
                opening_total,closing_total,delta_total,sample_count,source_json,created_at)
                VALUES(0,?,'user','1','Fixture user',0,?,?,2,?,?)''',
                (day, pages, pages, '{"device_identity":{"serial":"TEST","name":"Fixture"}}', day))
            conn.execute('''INSERT INTO daily_raw_counter_records(device_slot,record_date,canonical_field,label,
                subject_type,subject_id,subject_name,opening_value,closing_value,delta_value,device_serial_snapshot,created_at,updated_at)
                VALUES(0,?,'counter.print.full_color.a4','A4','user','1','Fixture user',0,?,?,'TEST',?,?)''',
                (day, pages, pages, day, day))
        post_daily_record(ledger_source)
    if archived:
        from services.report_detail_archive import ReportDetailArchive
        archive = ReportDetailArchive()
        batches = []
        while batch := archive.archive_batch(today=date(2028, 7, 5)):
            batches.append(batch)
        assert sum(batch['events'] for batch in batches) == 3
    stages = []
    target = tmp_path/'selected.xlsx'
    selected = date(2026, 7, 2)
    export_report_selection(target, selected, selected, database_path=database.get_db_path(),
        devices=(device,), cancelled=Event(), progress=lambda percent, text: stages.append(percent))
    book = load_workbook(target, data_only=True)
    assert book['มุมมองผู้ใช้']['AA5'].value == 3
    assert book['Legacy-ผู้ใช้']['B5'].value == 3
    assert book['มุมมองรายวัน']['A5'].value.date() == selected
    assert book['มุมมองรายวัน'].max_row == 5
    book.close()
    assert stages == sorted(stages) and stages[0] == 5 and stages[-1] == 99
    with pytest.raises(EmptyReportError):
        export_report_selection(tmp_path/'empty.xlsx', date(2020, 1, 1), date(2020, 1, 2),
            database_path=database.get_db_path(), cancelled=Event(), progress=lambda *_: None)
    assert not (tmp_path/'empty.xlsx').exists()


def test_page_date_cancel_does_not_read_database_and_export_busy_survives_refresh(monkeypatch):
    import v5_pages
    import services.report_export_job as job
    app = QApplication.instance() or QApplication([])
    page = v5_pages.PageHistoryV5(staged=True)
    page.report_center._period = ('2026-07-01', '2026-07-31')
    page.report_center._has_data = True
    monkeypatch.setattr(v5_pages.ReportExportPeriodDialog, 'get_range', lambda *a, **kw: None)
    monkeypatch.setattr(job, 'load_usage_history', lambda *a, **kw: pytest.fail('cancel read history'))
    try:
        page._export_report_center_excel()
        assert getattr(page, '_export_controller', None) is None
        page.report_center.set_export_busy(True)
        page.report_center._update_loading_visibility()
        assert not page.report_center.export_button.isEnabled()
        page.report_center.set_export_busy(False)
        assert page.report_center.export_button.isEnabled()
    finally:
        page.shutdown_report_loading()
        page.close()
        app.processEvents()


def test_page_export_stays_responsive_during_slow_database_read(monkeypatch, tmp_path):
    import v5_pages
    import services.report_export_job as job
    app = QApplication.instance() or QApplication([])
    page = v5_pages.PageHistoryV5(staged=True)
    page.report_center._period = ('2026-07-01', '2026-07-31')
    page.report_center._has_data = True
    page.report_center._update_loading_visibility()
    page.show()
    started, release = Event(), Event()
    messages, ticks = [], []
    ui_thread = get_ident()
    def slow_read(*args, **kwargs):
        assert get_ident() != ui_thread
        started.set()
        release.wait(3)
        raise OSError('synthetic failed database read')
    monkeypatch.setattr(job, 'load_usage_history', slow_read)
    monkeypatch.setattr(v5_pages.ReportExportPeriodDialog, 'get_range',
                        lambda *a, **kw: (date(2026, 7, 1), date(2026, 7, 31)))
    monkeypatch.setattr(v5_pages.QFileDialog, 'getSaveFileName', lambda *a, **kw: (str(tmp_path/'failed.xlsx'), ''))
    monkeypatch.setattr(v5_pages.QMessageBox, 'critical', lambda *a: messages.append(a[-1]))
    timer = QTimer(page)
    timer.setInterval(10)
    timer.timeout.connect(lambda: ticks.append(1))
    try:
        timer.start()
        page.report_center.export_button.click()
        wait_until(started.is_set)
        QTest.qWait(50)
        assert len(ticks) >= 2 and not page.report_center.export_button.isEnabled()
        page.report_center.tabs.setCurrentIndex(3)
        page._export_report_center_excel()  # Direct duplicate request must also be ignored.
        release.set()
        wait_until(lambda: not page._export_controller.busy)
        assert len(messages) == 1 and 'synthetic failed database read' in messages[0]
        assert page.report_center.export_button.isEnabled()
        assert not (tmp_path/'failed.xlsx').exists()
    finally:
        release.set()
        timer.stop()
        page.shutdown_report_loading()
        page.close()
        app.processEvents()
