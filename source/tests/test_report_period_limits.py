"""Public date limits, rejected-load isolation, and exact export handoff."""
from concurrent.futures import Future
from dataclasses import replace
from datetime import date, timedelta
from pathlib import Path
from threading import Event
from types import SimpleNamespace

import pytest
from PyQt6.QtCore import QDate
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication, QDialog, QPushButton
from openpyxl import load_workbook

import database
import v5_pages
from i18n import set_language
from services import report_snapshot
from services.report_period_policy import (
    EXPORT_YEARS, PREVIEW_YEARS, ReportPeriodLimitError,
    current_preview_period, latest_allowed_end, latest_export_start,
    validate_report_period,
)
from ui.cloud_reports import CloudReportsWidget
from test_report_custom_period import _usage
from excel_export_fixture import report_for


@pytest.mark.parametrize('start,years,end', [
    (date(2010, 1, 1), 1, date(2010, 12, 31)),
    (date(2024, 1, 1), 1, date(2024, 12, 31)),
    (date(2024, 2, 29), 1, date(2025, 2, 28)),
    (date(2023, 3, 1), 1, date(2024, 2, 29)),
    (date(2021, 9, 18), 5, date(2026, 9, 17)),
    (date(2024, 2, 29), 5, date(2029, 2, 28)),
])
def test_calendar_limits_allow_boundary_and_reject_one_extra_day(start, years, end):
    assert latest_allowed_end(start, years) == end
    validate_report_period(start, end, years=years)
    with pytest.raises(ReportPeriodLimitError) as error:
        validate_report_period(start, end + timedelta(days=1), years=years)
    assert error.value.latest_end == end
    assert end.strftime('%d/%m/%Y') in str(error.value)


def test_default_is_current_calendar_month_and_reverse_range_is_invalid():
    assert current_preview_period(date(2024, 2, 9)) == ('2024-02-01', '2024-02-29')
    with pytest.raises(ValueError):
        validate_report_period(date(2026, 2, 1), date(2026, 1, 31), years=PREVIEW_YEARS)
    for end in (date(2024, 2, 29), date(2025, 2, 28), date(2026, 12, 31)):
        validate_report_period(latest_export_start(end), end, years=EXPORT_YEARS)


@pytest.mark.parametrize('index', range(5))
def test_preview_worker_rejects_long_range_before_any_database_access(monkeypatch, index):
    monkeypatch.setattr(report_snapshot, 'report_read_database',
                        lambda *a, **kw: pytest.fail('Rejected Preview opened a database'))
    with pytest.raises(ReportPeriodLimitError):
        report_snapshot.load_report_snapshot(index, ('2021-01-01', '2025-12-31'),
            cache=None, subject_cache=None, repository=None, cancelled=Event())


def test_initial_preview_uses_current_month_even_when_only_old_history_exists(monkeypatch):
    expected = ('2026-09-01', '2026-09-30')
    monkeypatch.setattr(report_snapshot, 'current_preview_period', lambda: expected)
    monkeypatch.setattr(report_snapshot, 'report_slots', lambda: [0])
    old = _usage(date(2021, 1, 1), date(2021, 1, 31), 12)
    monkeypatch.setattr(report_snapshot, 'load_usage_history',
                        lambda *a, **kw: (old,) if kw['granularity'] == 'monthly' else ())
    snapshot = report_snapshot.load_report_snapshot(2, None, cache=None, subject_cache=None,
        repository=SimpleNamespace(list=lambda **kw: ()), cancelled=Event())
    assert snapshot.data['period'] == expected
    assert snapshot.data['records'] == ()


@pytest.mark.parametrize('language', ['th', 'en'])
def test_rejected_preview_keeps_dates_and_exports_exact_range_without_loading(monkeypatch, language):
    app = QApplication.instance() or QApplication([])
    set_language(language)
    page = v5_pages.PageHistoryV5(staged=True)
    cloud = page.report_center
    requests, dialogs = [], []
    def submit(*args, **kwargs):
        requests.append(args[2])
        return Future()
    def choose(_parent, start, end, **kwargs):
        dialogs.append((start, end))
        return None
    monkeypatch.setattr(page._snapshot_pool, 'submit', submit)
    monkeypatch.setattr(v5_pages.ReportExportPeriodDialog, 'get_selection', choose)
    monkeypatch.setattr(v5_pages.QFileDialog, 'getSaveFileName',
                        lambda *a, **kw: pytest.fail('Cancelled handoff opened Save'))
    page.resize(1280, 800)
    page.show()
    try:
        cloud.set_periods((), ('2026-01-01', '2026-01-31'))
        cloud.period_start_edit.setDate(QDate(2021, 1, 1))
        cloud.period_end_edit.setDate(QDate(2025, 12, 31))
        for index in range(5):
            cloud.tabs.setCurrentIndex(index)
            page.refresh_visible_cached()
        QTest.qWait(150)
        assert requests == []
        assert cloud.preview_limit_panel.isVisible()
        assert cloud.updated_label.isHidden()
        assert all(scroll.isHidden() for scroll in cloud.tab_scrolls)
        assert all(grid.parentWidget().height() < 150 for grid, _ in cloud._period_toolbars
                   if grid.parentWidget().isVisible())
        assert all(start.date() == QDate(2021, 1, 1) and end.date() == QDate(2025, 12, 31)
                   for start, end in cloud.period_edit_pairs)
        assert ('เกิน 12 เดือน' if language == 'th' else 'exceeds 12 months') in cloud.preview_limit_label.text()
        evidence = Path(__file__).resolve().parents[2] / 'outputs' / 'report-period-policy'
        evidence.mkdir(parents=True, exist_ok=True)
        page.grab().save(str(evidence / f'preview-limit-{language}.png'))
        cloud.preview_export_button.click()
        assert dialogs == [(date(2021, 1, 1), date(2025, 12, 31))]
        assert cloud.current_period() == ('2021-01-01', '2025-12-31')
        assert not getattr(page, '_export_controller', None)
        cloud.period_end_edit.setDate(QDate(2021, 12, 31))
        assert requests == [('2021-01-01', '2021-12-31')]
        assert cloud.preview_limit_panel.isHidden()
        assert not cloud.loading_panel.isHidden()  # The valid selection is now loading.
        assert all(not scroll.isHidden() for scroll in cloud.tab_scrolls)
    finally:
        page.shutdown_report_loading()
        page.close()
        app.processEvents()


def test_usage_filters_old_years_and_partial_months_without_changing_posted_totals():
    app = QApplication.instance() or QApplication([])
    widget = CloudReportsWidget()
    old = _usage(date(2021, 9, 1), date(2021, 9, 30), 999)
    monthly = _usage(date(2026, 9, 1), date(2026, 9, 30), 293, cost=42.75)
    daily = tuple(_usage(date(2026, 9, day), date(2026, 9, day), total)
                  for day, total in ((1, 75), (3, 10), (5, 8), (9, 200)))
    widget.set_data(records=(monthly,), monthly_records=(old, monthly), daily_records=daily,
                    subjects=(), devices=(), period=('2026-09-03', '2026-09-07'), last_sync_at=None)
    try:
        assert len(widget.period_edit_pairs) == 5  # Usage now has the shared selector too.
        for granularity in ('D', 'M', 'Y'):
            assert sum(row.total for row in widget._usage_records(granularity)) == 18
        widget.set_periods((), ('2026-09-01', '2026-09-30'))
        assert widget._usage_records('M') == (monthly,)
        assert widget._usage_records('Y')[0].net_cost == 42.75
        assert all(row.period_start.year == 2026 for row in widget._usage_records('D'))
    finally:
        widget.close()
        app.processEvents()


def test_late_preview_result_cannot_replace_rejected_range(monkeypatch):
    app = QApplication.instance() or QApplication([])
    page = v5_pages.PageHistoryV5(staged=True)
    cloud = page.report_center
    future = Future()
    future.set_running_or_notify_cancel()
    monkeypatch.setattr(page._snapshot_pool, 'submit', lambda *a, **kw: future)
    try:
        cloud.set_periods((), ('2026-01-01', '2026-01-31'))
        page.refresh_visible_cached()
        assert len(page._snapshot_jobs) == 1
        cancellation = next(iter(page._snapshot_jobs.values()))[1]
        cloud.period_start_edit.setDate(QDate(2021, 1, 1))
        assert cancellation.is_set()
        future.set_result(SimpleNamespace(data={'period': ('2026-01-01', '2026-01-31')}))
        page._collect_report_snapshots()
        assert cloud.current_period() == ('2021-01-01', '2026-01-31')
        assert cloud._preview_blocked_period == cloud.current_period()
        assert page._snapshot_cache == {}
        assert not cloud.has_data
    finally:
        page.shutdown_report_loading()
        page.close()
        app.processEvents()


@pytest.mark.parametrize('language', ['th', 'en'])
def test_export_dialog_allows_five_years_blocks_overflow_and_offers_latest_five(language):
    app = QApplication.instance() or QApplication([])
    set_language(language)
    dialog = v5_pages.ReportExportPeriodDialog(date(2022, 1, 1), date(2026, 12, 31),
        available_start=date(2010, 1, 1), available_end=date(2026, 12, 31))
    try:
        assert dialog.export_button.isEnabled()
        assert ('วัน' if language == 'th' else 'days') in dialog.status.text()
        dialog.show()
        QTest.qWait(100)
        evidence = Path(__file__).resolve().parents[2] / 'outputs' / 'report-period-policy'
        evidence.mkdir(parents=True, exist_ok=True)
        dialog.grab().save(str(evidence / f'export-five-years-{language}.png'))
        dialog.start_date.setDate(QDate(2021, 12, 31))
        assert not dialog.export_button.isEnabled()
        assert ('5 ปี' if language == 'th' else '5 years') in dialog.status.text()
        assert '30/12/2026' in dialog.status.text()
        dialog.accept()
        assert dialog.result() != QDialog.DialogCode.Accepted
        label = '5 ปีล่าสุด' if language == 'th' else 'Latest 5 years'
        next(button for button in dialog.findChildren(QPushButton) if button.text() == label).click()
        assert dialog.selected_range() == (date(2022, 1, 1), date(2026, 12, 31))
        assert dialog.export_button.isEnabled()
    finally:
        dialog.close()
        app.processEvents()


def test_export_worker_rejects_over_five_years_before_reads_and_keeps_existing_file(tmp_path, monkeypatch):
    from services import report_export_job as job
    target = tmp_path / 'existing.xlsx'
    target.write_bytes(b'existing workbook')
    monkeypatch.setattr(job, 'report_read_database',
                        lambda *a, **kw: pytest.fail('Oversized export opened a database'))
    with pytest.raises(ReportPeriodLimitError):
        job.export_report_selection(target, date(2021, 1, 1), date(2026, 1, 1),
            database_path=database.get_db_path(), cancelled=Event(), progress=lambda *_: None)
    assert target.read_bytes() == b'existing workbook'


def test_excel_writer_accepts_five_year_selection_and_preserves_it_when_overflow_rejected(tmp_path):
    from services.report_excel_export import export_cloud_reports_xlsx
    report = replace(report_for(users=2, days=1, printers=1),
                     period=('2022-01-01', '2026-12-31'))
    target = tmp_path / 'five-years.xlsx'
    export_cloud_reports_xlsx(target, report)
    original = target.read_bytes()
    with target.open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=True)
        assert len(list(book['มุมมองผู้ใช้'].iter_rows(min_row=5))) == 2
        book.close()
    with pytest.raises(ReportPeriodLimitError):
        export_cloud_reports_xlsx(target, replace(report, period=('2021-12-31', '2026-12-31')))
    assert target.read_bytes() == original
    assert not list(tmp_path.glob('.komi-export-*'))
