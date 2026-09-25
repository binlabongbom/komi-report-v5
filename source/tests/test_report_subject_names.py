from dataclasses import replace
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from services.report_excel_export import CloudReportsExportData, DailyCategoryCost, export_cloud_reports_xlsx
from services.report_models import build_user_rows
from services.report_user_counters import UserCounter
from test_report_excel_export import _raw, _subject


def subject(day=1, *, name="00123", slot=0, serial="A", subject_id="00123", usage=5):
    observed = date(2026, 8, day)
    return replace(_subject(observed, slot=slot, identity=f"serial:{serial}",
                            name=f"Printer {serial}", usage=usage),
                   subject_id=subject_id, subject_name=name,
                   latest_at=f"{observed.isoformat()}T09:00:00+07:00")


def raw(day=1, *, name="00123", slot=0, serial="A", subject_id="00123", usage=5):
    observed = date(2026, 8, day)
    return replace(_raw(observed, slot=slot, field="counter.print.full_color.a4",
                        label="Print A4 color", delta=usage, subject=True, serial=serial),
                   subject_id=subject_id, subject_name=name,
                   last_success=f"{observed.isoformat()}T09:00:00+07:00")


def excel_rows(tmp_path, subjects, records=(), *, filename="names.xlsx", costs=()):
    report = CloudReportsExportData((), (), (), (), tuple(subjects), (),
        ("2026-08-01", "2026-08-31"), None, datetime(2026, 8, 31, tzinfo=UTC),
        raw_counter_records=tuple(records), category_costs=tuple(costs), sections=("users",))
    workbook = load_workbook(export_cloud_reports_xlsx(tmp_path / filename, report), data_only=True)
    try:
        return tuple(workbook["มุมมองผู้ใช้"].iter_rows(min_row=5, values_only=True))
    finally:
        workbook.close()


@pytest.mark.parametrize("placeholder", ["00123", "123", ""])
def test_report_name_survives_later_placeholder_and_preserves_usage(tmp_path, placeholder):
    subjects = (subject(name="Alice"), subject(2, name=placeholder, usage=7))
    records = (raw(name="Alice"), raw(2, name=placeholder, usage=7))
    rows = build_user_rows(subjects, records)
    assert len(rows) == 1
    assert (rows[0]["name"], rows[0]["usage"], rows[0]["print"]) == ("Alice", 12, 12)
    exported = excel_rows(tmp_path, subjects, records)
    assert len(exported) == 1
    assert exported[0][:3] == ("Alice", "user", "00123")
    assert exported[0][3] == exported[0][-2] == 12
    assert subjects[-1].subject_name == placeholder


def test_report_uses_latest_real_rename_even_with_reverse_input_order(tmp_path):
    subjects = (subject(3, name="Bob", usage=7), subject(name="Alice"))
    records = (raw(3, name="Bob", usage=7), raw(name="Alice"))
    rows = build_user_rows(subjects, records)
    assert [(row["name"], row["usage"], row["print"]) for row in rows] == [("Bob", 12, 12)]
    assert excel_rows(tmp_path, subjects, records)[0][0] == "Bob"


def test_raw_real_name_fills_subject_placeholder_and_keeps_latest_real_name(tmp_path):
    subjects = (subject(4, name="123", usage=15),)
    records = (raw(3, name="123"), raw(2, name="Bob"), raw(name="Alice"))
    rows = build_user_rows(subjects, records)
    assert [(row["name"], row["usage"], row["print"]) for row in rows] == [("Bob", 15, 15)]
    assert excel_rows(tmp_path, subjects, records)[0][0] == "Bob"


def test_subject_counter_date_does_not_mask_a_later_raw_real_rename(tmp_path):
    subjects = (replace(subject(4, name="Alice"), name_observed_on=date(2026, 8, 1)),)
    records = (raw(3, name="Bob"),)
    assert build_user_rows(subjects, records)[0]["name"] == "Bob"
    assert excel_rows(tmp_path, subjects, records)[0][0] == "Bob"


def test_numeric_placeholders_never_merge_across_slots_or_replaced_printers(tmp_path):
    subjects = (subject(name="123"), subject(slot=1, serial="B", usage=7),
                subject(name="", serial="C", usage=9))
    records = (raw(name="123"), raw(slot=1, serial="B", usage=7), raw(name="", serial="C", usage=9))
    rows = build_user_rows(subjects, records)
    assert len(rows) == len({row["key"] for row in rows}) == 3
    assert {row["name"] for row in rows} == {"00123"}
    assert sorted((row["usage"], row["print"]) for row in rows) == [(5, 5), (7, 7), (9, 9)]
    exported = excel_rows(tmp_path, subjects, records)
    assert len(exported) == 3
    assert sorted(row[-2] for row in exported) == [5, 7, 9]
    assert [sum(value or 0 for value in row[3:-2:25]) for row in exported] == [row[-2] for row in exported]


@pytest.mark.parametrize("foreign", [
    {"serial": "OLD"}, {"slot": 1, "serial": "B"}, {"subject_id": "123"},
])
def test_other_identity_cannot_supply_name_or_raw_amounts(tmp_path, foreign):
    subjects = (subject(name="", usage=7),)
    records = (raw(name="123", usage=7), raw(name="Foreign name", usage=100, **foreign),
               replace(raw(name="Account name", usage=200), subject_type="account"))
    rows = build_user_rows(subjects, records)
    assert [(row["name"], row["usage"], row["print"]) for row in rows] == [("00123", 7, 7)]
    exported = excel_rows(tmp_path, subjects, records)
    assert len(exported) == 1
    assert (exported[0][0], exported[0][3], exported[0][-2]) == ("00123", 7, 7)


def test_same_name_groups_across_printers_but_cannot_name_another_slots_placeholder(tmp_path):
    first = subject(name="Alice")
    other = subject(slot=1, serial="B", name="Alice", usage=7)
    for named, expected_count in ((other, 1), (replace(other, subject_name="123"), 2)):
        subjects = (first, named)
        rows = build_user_rows(subjects, ())
        exported = excel_rows(tmp_path, subjects)
        assert len(rows) == len(exported) == expected_count
        assert sum(row["usage"] for row in rows) == sum(row[-2] for row in exported) == 12
        assert sorted(row["name"] for row in rows) == sorted(row[0] for row in exported)


@pytest.mark.parametrize("printer_count", [2, 3])
def test_same_person_with_different_printer_ids_has_one_row_and_separate_slot_totals(tmp_path, printer_count):
    people = ((0, "A", "00123", "บอม", 5, 2),
              (1, "B", "00456", " บอม ", 7, 3),
              (2, "C", "00789", "บอม", 11, 4))[:printer_count]
    subjects = tuple(subject(slot=slot, serial=serial, subject_id=identity, name=name, usage=usage)
                     for slot, serial, identity, name, usage, _rate in people)
    subject_raw = tuple(raw(slot=slot, serial=serial, subject_id=identity, name=name, usage=usage)
                        for slot, serial, identity, name, usage, _rate in people)
    records = subject_raw + tuple(replace(item, subject_type="", subject_id="", subject_name="")
                                  for item in subject_raw)
    costs = tuple(DailyCategoryCost(slot, f"serial:{serial}", date(2026, 8, 1), a4_color=usage * rate)
                  for slot, serial, _identity, _name, usage, rate in people)
    rows = build_user_rows(subjects, records, category_costs=costs)
    exported = excel_rows(tmp_path, subjects, records, costs=costs)
    total = sum(item[4] for item in people)
    assert len(rows) == len(exported) == 1
    assert (rows[0]["name"], rows[0]["usage"], rows[0]["print"]) == ("บอม", total, total)
    assert exported[0][:3] == ("บอม", "user", " | ".join(
        f"Slot {slot + 1}: {identity}" for slot, _serial, identity, _name, _usage, _rate in people))
    assert exported[0][-2] == total
    for slot, _serial, _identity, _name, usage, rate in people:
        assert rows[0]["by_slot"][slot]["total"] == usage
        assert rows[0]["by_slot"][slot]["cost"] == usage * rate
        assert rows[0]["by_slot"][slot]["cost_available"]
        start = 3 + slot * 25
        assert [exported[0][start + offset] for offset in (0, 1, 21, 23, 24)] == [
            usage, usage * rate, usage, usage, usage * rate]


def test_same_name_multiple_local_ids_do_not_duplicate_raw_totals(tmp_path):
    subjects = (subject(name="บอม"), subject(name="บอม", subject_id="00456", usage=7))
    records = (raw(name="บอม"), raw(name="บอม", subject_id="00456", usage=7))
    rows = build_user_rows(subjects, records)
    exported = excel_rows(tmp_path, subjects, records)
    assert len(rows) == len(exported) == 1
    assert (rows[0]["usage"], rows[0]["print"]) == (12, 12)
    assert exported[0][2] == "Slot 1: 00123 | Slot 1: 00456"
    assert exported[0][3] == exported[0][-2] == 12


def test_same_local_id_with_different_real_names_stays_separate(tmp_path):
    subjects = (subject(name="บอม"), subject(slot=1, serial="B", name="วิภา", usage=7))
    records = (raw(name="บอม"), raw(slot=1, serial="B", name="วิภา", usage=7))
    rows = build_user_rows(subjects, records)
    exported = excel_rows(tmp_path, subjects, records)
    assert len(rows) == len(exported) == 2
    assert {(row["name"], row["usage"]) for row in rows} == {("บอม", 5), ("วิภา", 7)}
    assert {(row[0], row[-2]) for row in exported} == {("บอม", 5), ("วิภา", 7)}


def test_slot_only_legacy_subject_uses_only_unambiguous_physical_identity(tmp_path):
    subjects = (replace(subject(name=""), device_identity="", device_serial_snapshot=""),)
    records = (raw(name="Alice"),)
    rows = build_user_rows(subjects, records)
    assert [(row["name"], row["print"]) for row in rows] == [("Alice", 5)]
    assert excel_rows(tmp_path, subjects, records)[0][0:4] == ("Alice", "user", "00123", 5)
    ambiguous = records + (raw(serial="OLD", name="Someone else", usage=100),)
    rows = build_user_rows(subjects, ambiguous)
    assert [(row["name"], row["print"]) for row in rows] == [("00123", 0)]
    exported = excel_rows(tmp_path, subjects, ambiguous)
    assert (exported[0][0], exported[0][3], exported[0][-2]) == ("00123", None, 5)


def test_compact_and_full_raw_choose_the_same_latest_name():
    subjects = (subject(4, name="123", usage=15),)
    records = (raw(3, name="123"), raw(2, name="Bob"), raw(name="Alice"))
    compact = tuple(UserCounter(*(getattr(record, field) for field in UserCounter._fields))
                    for record in records)
    assert build_user_rows(subjects, compact) == build_user_rows(subjects, records)


@pytest.mark.parametrize("include_slot_cost", [False, True])
def test_mixed_legacy_and_serial_subjects_count_each_raw_fact_and_cost_once(tmp_path, include_slot_cost):
    # Distinct facts on the same day can retain old slot-only metadata alongside
    # a known serial. Name compatibility must not repeat Raw in two columns.
    subjects = (replace(subject(name=""), device_identity="", device_serial_snapshot=""),
                subject(name="Alice", usage=7))
    records = (
        replace(raw(name=""), device_serial_snapshot=""),
        raw(name="Alice", usage=7),
        replace(raw(usage=12), subject_type="", subject_id="", subject_name=""),
    )
    cost = DailyCategoryCost(0, "serial:A", date(2026, 8, 1), a4_color=24)
    costs = (cost, replace(cost, device_identity="slot:0", a4_color=999)) if include_slot_cost else (cost,)
    rows = build_user_rows(subjects, records, category_costs=costs)
    assert len(rows) == 1
    assert (rows[0]["name"], rows[0]["usage"], rows[0]["print"]) == ("Alice", 12, 12)
    assert rows[0]["by_slot"][0]["cost"] == 24
    assert rows[0]["by_slot"][0]["cost_available"]
    exported = excel_rows(tmp_path, subjects, records, costs=costs)
    assert len(exported) == 1
    assert len(exported[0]) == 30  # One physical printer, with 25 detail columns.
    assert exported[0][0] == "Alice"
    assert [exported[0][index] for index in (3, 4, 24, 26, 27, 28)] == [12, 24, 12, 12, 24, 12]


def test_legacy_and_serial_days_share_one_column_and_retain_available_daily_prices(tmp_path):
    subjects = (replace(subject(name="Alice"), device_identity="", device_serial_snapshot=""),
                subject(2, name="Alice"))
    subject_raw = (raw(name="Alice"), raw(2, name="Alice"))
    device_raw = tuple(replace(item, subject_type="", subject_id="", subject_name="")
                       for item in subject_raw)
    records = subject_raw + device_raw
    costs = (DailyCategoryCost(0, "slot:0", date(2026, 8, 1), a4_color=10),
             DailyCategoryCost(0, "serial:A", date(2026, 8, 2), a4_color=15))
    rows = build_user_rows(subjects, records, category_costs=costs)
    assert len(rows) == 1
    assert (rows[0]["usage"], rows[0]["print"]) == (10, 10)
    assert rows[0]["by_slot"][0]["cost"] == 25
    assert rows[0]["by_slot"][0]["cost_available"]
    exported = excel_rows(tmp_path, subjects, records, costs=costs)
    assert len(exported) == 1 and len(exported[0]) == 30
    assert [exported[0][index] for index in (3, 4, 24, 26, 27, 28)] == [10, 25, 10, 10, 25, 10]


def test_competing_serials_keep_columns_and_cannot_inherit_slot_price(tmp_path):
    subjects = (subject(name="Alice"), subject(serial="B", name="Alice", usage=7))
    subject_raw = (raw(name="Alice"), raw(serial="B", name="Alice", usage=7))
    device_raw = tuple(replace(item, subject_type="", subject_id="", subject_name="")
                       for item in subject_raw)
    records = subject_raw + device_raw
    costs = (DailyCategoryCost(0, "slot:0", date(2026, 8, 1), a4_color=24),)
    rows = build_user_rows(subjects, records, category_costs=costs)
    assert (rows[0]["usage"], rows[0]["print"]) == (12, 12)
    assert not rows[0]["by_slot"][0]["cost_available"]
    exported = excel_rows(tmp_path, subjects, records, costs=costs)
    assert len(exported) == 1 and len(exported[0]) == 55
    assert [exported[0][index] for index in (3, 28, 53)] == [5, 7, 12]
    assert exported[0][27] is exported[0][52] is None


def test_name_recovery_leaves_existing_user_grouping_pages_and_prices_unchanged(tmp_path):
    from excel_export_fixture import report_for

    report = report_for(users=2, days=2, printers=2)
    subjects = tuple(replace(item, subject_name=item.subject_id) for item in report.subjects)
    records = tuple(replace(item, subject_name=item.subject_id)
                    if item.is_subject and item.record_date.day == 2 else item
                    for item in report.raw_counter_records)
    expected = build_user_rows(report.subjects, report.raw_counter_records, category_costs=report.category_costs)
    assert build_user_rows(subjects, records, category_costs=report.category_costs) == expected
    output_values = []
    for label, payload in (("baseline", report), ("recovered", replace(report, subjects=subjects,
                                                                       raw_counter_records=records))):
        book = load_workbook(export_cloud_reports_xlsx(tmp_path / f"{label}.xlsx",
                             replace(payload, sections=("users",))), data_only=True)
        try:
            output_values.append(tuple(book["มุมมองผู้ใช้"].values))
        finally:
            book.close()
    assert output_values[0] == output_values[1]


def test_same_name_local_ids_keep_independent_counter_format_derivation(tmp_path):
    subjects = (subject(name="บอม", usage=5),
                subject(name="บอม", subject_id="00456", usage=7))
    records = (
        replace(raw(name="บอม", usage=5),
                canonical_field="counter.print.full_color_total", label="Full Color Total"),
        raw(name="บอม", subject_id="00456", usage=7),
    )
    rows = build_user_rows(subjects, records)
    exported = excel_rows(tmp_path, subjects, records)
    assert len(rows) == len(exported) == 1
    assert (rows[0]["name"], rows[0]["usage"], rows[0]["print"]) == ("บอม", 12, 12)
    assert rows[0]["by_slot"][0]["print_a4_color"] == 12
    assert exported[0][2] == "Slot 1: 00123 | Slot 1: 00456"
    assert [exported[0][index] for index in (3, 24, 26, 28)] == [12, 12, 12, 12]


def test_same_name_local_ids_preserve_individually_rounded_costs(tmp_path):
    subjects = (subject(name="บอม", usage=1),
                subject(name="บอม", subject_id="00456", usage=1))
    records = (
        raw(name="บอม", usage=1),
        raw(name="บอม", subject_id="00456", usage=1),
        replace(raw(usage=3), subject_type="", subject_id="", subject_name=""),
    )
    costs = (DailyCategoryCost(0, "serial:A", date(2026, 8, 1), a4_color=1),)
    rows = build_user_rows(subjects, records, category_costs=costs)
    exported = excel_rows(tmp_path, subjects, records, costs=costs)
    assert len(rows) == len(exported) == 1
    assert (rows[0]["usage"], rows[0]["print"]) == (2, 2)
    # Each exact local account owns 1/3 of the cost, rounded to 0.33 first.
    assert rows[0]["by_slot"][0]["cost"] == 0.66
    assert rows[0]["by_slot"][0]["cost_available"]
    assert exported[0][2] == "Slot 1: 00123 | Slot 1: 00456"
    assert [exported[0][index] for index in (3, 4, 24, 26, 27, 28)] == [
        2, 0.66, 2, 2, 0.66, 2]


def test_merged_positive_usage_without_raw_keeps_total_cost_unavailable(tmp_path):
    subjects = (subject(name="บอม", usage=1),
                subject(name="บอม", subject_id="00456", usage=1))
    records = (raw(name="บอม", usage=1),
               replace(raw(usage=3), subject_type="", subject_id="", subject_name=""))
    costs = (DailyCategoryCost(0, "serial:A", date(2026, 8, 1), a4_color=1),)
    rows = build_user_rows(subjects, records, category_costs=costs)
    exported = excel_rows(tmp_path, subjects, records, costs=costs)
    assert len(rows) == len(exported) == 1
    assert rows[0]["usage"] == exported[0][26] == exported[0][-2] == 2
    assert not rows[0]["by_slot"][0]["cost_available"]
    assert exported[0][27] is None
