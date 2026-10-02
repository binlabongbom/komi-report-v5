"""Verify the requested chart semantics, compact controls and selective export."""
from dataclasses import replace
from datetime import date
from pathlib import Path
from threading import Event

import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication

from excel_export_fixture import report_for
from i18n import apply_language, set_language
from services.error_codes import ERROR_CODES
from services.error_reference import CHECKS
from services.report_chart_models import FIVE_CATEGORIES
from services.report_excel_export import export_cloud_reports_xlsx
from services.report_export_options import ALL_REPORT_SECTIONS, REPORT_SECTIONS
from services.support_diagnostics import diagnose
from ui.cloud_reports import CloudReportsWidget
from ui.error_list import ErrorListPanel
from ui.reports.export_dialog import ReportExportPeriodDialog


@pytest.fixture
def report_widget():
    report = report_for(users=2, days=3, printers=3)
    report = replace(report, devices=tuple(replace(device, name=f'KM bizhub C{450 + slot * 100}i Test printer')
        for slot, device in enumerate(report.devices)))
    widget = CloudReportsWidget()
    widget.resize(1280, 850)
    widget.set_periods((report.period,), report.period)
    widget.set_data(records=report.selected_records, daily_records=report.selected_daily_records,
        usage_daily_records=report.usage_daily_records, monthly_records=report.selected_records,
        subjects=report.subjects, devices=report.devices, period=report.period,
        last_sync_at=None, raw_counter_records=report.raw_counter_records)
    yield widget
    widget.close()


@pytest.mark.parametrize('language', ['th', 'en'])
def test_overview_and_printers_five_bars_share_paper_semantics(report_widget, language):
    widget = report_widget
    set_language(language)
    apply_language(widget)
    for tab, view in ((0, widget.overview_mix), (2, widget.printer_mix)):
        widget.tabs.setCurrentIndex(tab)
        chart = view.chart()
        assert chart.axes(Qt.Orientation.Horizontal)[0].categories() == list(FIVE_CATEGORIES)
        bars = chart.series()[0].barSets()
        assert [bar.label() for bar in bars] == ['Print A4', 'Print A3', 'Copy A4', 'Copy A3', 'Scan']
        assert [sum(bar.at(i) for bar in bars) for i in range(5)] == [36, 36, 36, 36, 0]
        assert bars[0].color().lightness() < bars[1].color().lightness()
        assert bars[2].color().lightness() < bars[3].color().lightness()
    timeline = widget.overview_trend.chart()
    assert timeline.axes(Qt.Orientation.Horizontal)[0].categoriesLabels() == ['01/07', '02/07', '03/07']
    bars = timeline.series()[0].barSets()
    assert [sum(bar.at(i) for bar in bars) for i in range(6)] == [24] * 6
    for i in range(6):
        assert sum(bar.at(i) for bar in bars[(2 if i % 2 == 0 else 0): (4 if i % 2 == 0 else 2)]) == 0


def test_wrapped_legend_and_inline_filters_render(report_widget):
    widget = report_widget
    widget.show()
    widget.tabs.setCurrentIndex(2)
    QTest.qWait(150)
    legend = widget.printer_trend_legend
    items = [legend.layout().itemAt(i).widget() for i in range(legend.layout().count())]
    assert len(items) == 3
    assert all(f'Slot {i + 1}' in item.text() for i, item in enumerate(items))
    assert all(item.geometry().bottom() < legend.height() for item in items)
    assert legend.y() >= widget.printer_trend.geometry().bottom() + 8
    assert all(item.height() <= 26 for item in items)
    assert widget.printer_trend.parentWidget().height() == widget.printer_mix.parentWidget().height()
    assert widget.printers_table.parentWidget().y() >= widget.printer_trend.parentWidget().geometry().bottom() + 8
    out = Path(__file__).resolve().parents[2] / 'outputs' / 'ui-batch'
    out.mkdir(parents=True, exist_ok=True)
    widget.grab().save(str(out / 'printers.png'))
    widget.tabs.setCurrentIndex(0)
    QTest.qWait(100)
    widget.grab().save(str(out / 'overview.png'))
    widget.tabs.setCurrentIndex(3)
    QTest.qWait(100)
    first = widget.usage_device_filters[0]
    assert abs(widget.usage_filter_title.y() - first.y()) <= 4
    widget.grab().save(str(out / 'usage.png'))


def test_export_dialog_rejects_empty_selection_and_invalid_dates():
    dialog = ReportExportPeriodDialog(date(2026, 7, 1), date(2026, 7, 31))
    assert dialog.selected_sections() == ALL_REPORT_SECTIONS
    dialog._select_all(False)
    assert not dialog.export_button.isEnabled()
    dialog.section_checks['daily'].setChecked(True)
    assert dialog.export_button.isEnabled() and dialog.selected_sections() == ('daily',)
    dialog.start_date.setDate(QDate(2026, 8, 1))
    assert not dialog.export_button.isEnabled()
    dialog.start_date.setDate(QDate(2026, 7, 1))
    dialog.show()
    QTest.qWait(60)
    assert dialog.start_date.width() < dialog.width() / 2
    out = Path(__file__).resolve().parents[2] / 'outputs' / 'ui-batch'
    out.mkdir(parents=True, exist_ok=True)
    dialog.grab().save(str(out / 'export.png'))
    dialog.close()


@pytest.mark.parametrize('section,title,_english', REPORT_SECTIONS)
def test_excel_writes_only_the_selected_report(tmp_path, section, title, _english):
    report = replace(report_for(users=2, days=2, printers=2), sections=(section,))
    target = export_cloud_reports_xlsx(tmp_path / 'selected.xlsx', report)
    book = load_workbook(target, data_only=False)
    assert book.sheetnames == [title]
    assert book[title].max_row > 3
    book.close()


def test_invalid_selection_preserves_existing_file(tmp_path):
    target = tmp_path / 'existing.xlsx'
    target.write_bytes(b'original')
    with pytest.raises(ValueError):
        export_cloud_reports_xlsx(target, replace(report_for(users=1, days=1, printers=1), sections=()))
    assert target.read_bytes() == b'original'


def test_all_errors_are_searchable_and_copyable():
    assert CHECKS.keys() == ERROR_CODES.keys()
    panel = ErrorListPanel()
    panel.resize(1280, 700)
    panel.show()
    QTest.qWait(80)
    assert panel.table.rowCount() == len(ERROR_CODES)
    panel.search.setText('KOMI-AUTH-003')
    assert panel.table.rowCount() == 1
    panel.table.selectRow(0)
    panel.copy_button.click()
    assert 'KOMI-AUTH-003' in QApplication.clipboard().text()
    assert 'Password' in QApplication.clipboard().text()
    panel.search.setText('PKI')
    assert panel.table.rowCount() >= 2
    panel.search.setText('no-such-code')
    assert panel.table.rowCount() == 0 and not panel.copy_button.isEnabled()
    panel.search.setText('TLS')
    QTest.qWait(100)
    out = Path(__file__).resolve().parents[2] / 'outputs' / 'ui-batch'
    out.mkdir(parents=True, exist_ok=True)
    panel.grab().save(str(out / 'errors.png'))
    panel.close()
    assert diagnose(ModuleNotFoundError("No module named example", name='example')).code == 'dependency_missing'
    from services.report_export_job import EmptyReportError
    assert diagnose(EmptyReportError('test period')).error_code == 'KOMI-DATA-003'


def test_daily_export_does_not_read_users_or_raw(monkeypatch, tmp_path):
    import database
    from services import report_export_job as job
    report = report_for(users=1, days=1, printers=1)
    monkeypatch.setattr(job, 'load_usage_history', lambda *args, **kw: report.selected_daily_records)
    monkeypatch.setattr(job, 'load_subject_history', lambda *args, **kw: pytest.fail('unselected users read'))
    monkeypatch.setattr(job, 'load_raw_counter_history', lambda *args, **kw: pytest.fail('unselected raw read'))
    target = job.export_report_selection(tmp_path / 'daily.xlsx', date(2026, 7, 1), date(2026, 7, 1),
        database_path=database.get_db_path(), devices=report.devices, sections=('daily',),
        cancelled=Event(), progress=lambda *_: None)
    book = load_workbook(target)
    assert book.sheetnames == ['มุมมองรายวัน']
    book.close()


def test_staged_startup_resizes_before_lazy_tabs_are_created():
    widget = CloudReportsWidget(staged=True)
    widget.resize(1000, 700)
    widget._apply_responsive_layout()
    widget.show()
    QTest.qWait(80)
    widget.tabs.setCurrentIndex(2)
    widget._apply_responsive_layout()
    widget.close()
