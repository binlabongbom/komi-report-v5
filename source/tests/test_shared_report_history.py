from dataclasses import replace
from datetime import UTC, date, datetime, timedelta, timezone
from itertools import permutations
import hashlib
import json
import sqlite3
from threading import Event
import uuid

import pytest

import database
from services.device_config import DeviceConfig, DeviceConfigRepository
from services.monthly_history import load_usage_history
from services.raw_counter_history import RawCounterHistoryStore, load_raw_counter_history
from services.shared_report_history import load_shared_counter_days
from services.shared_report_index import backfill_batch
from services.slot_epoch_store import SlotEpochStore
from services.subject_history import SubjectDailyCacheStore, load_subject_history
from services.web_observations import WebObservationStore, _json
from test_subject_history import _reading


def reading(value, at):
    r = _reading(value, collected_at=at)
    device = replace(r.counters[0], canonical_field='counter.total.total',
                     subject_type=None, subject_id=None, subject_name=None, source_path='device.total')
    return replace(r, counters=(*r.counters, device))


def setup(at=None, value=100):
    at = at or datetime(2026, 9, 18, 17, tzinfo=UTC)  # Friday
    DeviceConfigRepository().replace_all([DeviceConfig(0, 'Shared printer', '192.0.2.10', serial='REPORT-A')])
    event = WebObservationStore().record(0, reading(value, at), at=at, serial='REPORT-A', printer_name='Shared printer')
    store = SlotEpochStore()
    epoch = store.confirm(store.preview(0, event['id'], operation='initial', actor='A', reason='Reviewed opening'))
    return event, epoch


def event_at(anchor, epoch, value, at):
    p = dict(anchor['payload'], format='komi.web.observation.v2', meter_epoch_id=epoch,
             read_id=str(uuid.uuid4()), observed_at=at.isoformat(), reading=reading(value, at).to_dict())
    p['reading']['counters'] = list(p['reading']['counters'])
    p['reading']['warnings'] = list(p['reading']['warnings'])
    return {'id': hashlib.sha256(_json(p).encode()).hexdigest(), 'payload': p}


def receive(event, epoch, path=None):
    return WebObservationStore(path).receive([event], expected_slot=0, expected_serial='REPORT-A', expected_epoch=epoch)


def shared(*, graph=False, start=None, end=None):
    with database._open(commit=False) as conn:
        return load_shared_counter_days(conn, [0], kind='device', graph=graph, start=start, end=end)


@pytest.mark.parametrize('order', list(permutations(range(3))))
def test_reordered_duplicates_preserve_real_manual_reads_and_one_total(order):
    anchor, epoch = setup(datetime(2026, 9, 18, 10, tzinfo=UTC))
    events = [event_at(anchor, epoch, v, datetime(2026, 9, 18, 10, minute, tzinfo=UTC))
              for v, minute in [(110, 10), (125, 15), (150, 30)]]
    for i in order:
        receive(events[i], epoch)
        receive(events[i], epoch)
    rows = shared()
    assert sum(r.delta for r in rows) == 50
    assert not {flag for r in rows for flag in r.issues}
    users = load_subject_history([0], today=date(2026, 9, 19))
    assert len(users) == 1 and users[0].period_usage == 50 and users[0].sample_count == 4
    with database._open(commit=False) as conn:
        assert conn.execute('SELECT COUNT(*) FROM web_observations').fetchone()[0] == 4
        assert conn.execute('SELECT COUNT(*) FROM shared_report_events').fetchone()[0] == 4
        assert conn.execute('SELECT COUNT(*) FROM shared_report_days').fetchone()[0] == 1
        assert conn.execute('SELECT COUNT(*) FROM billing_ledger').fetchone()[0] == 0


def test_weekend_allocation_and_fractional_graph_conserve_usage_without_fake_reads():
    anchor, epoch = setup()
    monday = event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC))
    receive(monday, epoch)
    rows = shared()
    sunday = next(r for r in rows if r.day == date(2026, 9, 20))
    assert sunday.delta == 301
    assert sunday.count == 0 and sunday.opening is None and sunday.last_at is None
    graph = {r.day: r for r in shared(graph=True)}
    assert graph[date(2026, 9, 19)].delta == graph[date(2026, 9, 20)].delta == 150.5
    assert graph[date(2026, 9, 19)].estimated
    assert sum(r.delta for r in graph.values()) == sum(r.delta for r in rows) == 301
    raw = load_raw_counter_history([0], period_start=date(2026, 9, 20), period_end=date(2026, 9, 20))
    device = next(r for r in raw if not r.subject_type)
    assert device.delta_value == 301 and device.closing_value is None
    assert any('ยอดสะสมช่วงขาดการอ่าน' in note for note in device.remarks)
    usage = load_usage_history([0], granularity='daily', today=date(2026, 9, 22))
    assert sum(r.provisional_total for r in usage) == 301
    assert all(r.total == r.net_cost == 0 for r in usage)


def test_late_historical_read_changes_allocation_not_total_and_keeps_latest_live():
    anchor, epoch = setup()
    receive(event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC)), epoch)
    receive(event_at(anchor, epoch, 200, datetime(2026, 9, 19, 11, tzinfo=UTC)), epoch)
    rows = {r.day: r for r in shared()}
    assert rows[date(2026, 9, 19)].delta == 100
    assert rows[date(2026, 9, 20)].delta == 201
    assert sum(r.delta for r in rows.values()) == 301
    assert not {flag for r in rows.values() for flag in r.issues}


def test_closed_legacy_resumed_row_blocks_duplicate_gap_even_outside_selected_range(tmp_path):
    anchor, epoch = setup()
    monday_at = datetime(2026, 9, 21, 9, tzinfo=UTC)
    receive(event_at(anchor, epoch, 401, monday_at), epoch)
    # Existing closed daily detail already accounts for this gap on Monday.
    raw = RawCounterHistoryStore()
    raw.update_reading(0, reading(100, monday_at), at=monday_at)
    raw.update_reading(0, reading(401, monday_at + timedelta(minutes=1)), at=monday_at + timedelta(minutes=1))
    subjects = SubjectDailyCacheStore(tmp_path / 'subject-cache')
    subjects.update_reading(0, reading(100, monday_at), at=monday_at)
    subjects.update_reading(0, reading(401, monday_at + timedelta(minutes=1)), at=monday_at + timedelta(minutes=1))
    subjects.finalize(0, monday_at.date())
    assert not load_raw_counter_history([0], period_start=date(2026, 9, 20), period_end=date(2026, 9, 20))
    assert not load_subject_history([0], period_start=date(2026, 9, 20), period_end=date(2026, 9, 20))
    with database._open(commit=False) as conn:
        assert conn.execute("SELECT delta_total FROM daily_subject_records WHERE record_date='2026-09-21'").fetchone()[0] == 301


def test_index_rebuild_is_batched_idempotent_and_never_opens_old_archives(monkeypatch):
    anchor, epoch = setup()
    receive(event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC)), epoch)
    before = shared()
    with database._open() as conn:
        conn.execute('DELETE FROM shared_report_events')
        conn.execute('DELETE FROM shared_report_days')
        conn.execute('DELETE FROM shared_report_months')
    assert backfill_batch(database.get_db_path(), limit=1) == 1
    assert backfill_batch(database.get_db_path(), limit=1) == 1
    assert backfill_batch(database.get_db_path(), limit=1) == 0
    assert shared() == before
    from services.web_observation_archive import WebObservationArchive
    while WebObservationArchive().archive_batch(today=date(2028, 1, 1)):
        pass
    monkeypatch.setattr(WebObservationStore, '_document', lambda *_: pytest.fail('Opened an archive during normal reports'))
    assert shared() == before
    assert backfill_batch(database.get_db_path()) == 0
    assert load_subject_history([0])[0].period_usage == 301


def test_intermediate_decrease_is_not_hidden_by_daily_extrema():
    anchor, epoch = setup(datetime(2026, 9, 18, 10, tzinfo=UTC))
    for hour, value in [(13, 200), (12, 120), (11, 150)]:
        receive(event_at(anchor, epoch, value, datetime(2026, 9, 18, hour, tzinfo=UTC)), epoch)
    rows = shared()
    assert sum(r.delta for r in rows) == 100
    assert 'lower_counter_or_clock_conflict' in rows[0].issues
    usage = load_usage_history([0], today=date(2026, 9, 19))
    assert not usage[0].provisional_cost_available


def test_epoch_calendar_is_shared_even_if_sender_uses_another_offset():
    zone = timezone(timedelta(hours=7))
    anchor, epoch = setup(datetime(2026, 9, 18, 23, 50, tzinfo=zone))
    receive(event_at(anchor, epoch, 110, datetime(2026, 9, 18, 17, 10, tzinfo=UTC)), epoch)
    rows = {r.day: r for r in shared()}
    assert rows[date(2026, 9, 19)].delta == 10


def test_report_snapshot_and_excel_input_use_received_subjects_without_local_cache():
    from services.report_snapshot import load_report_snapshot
    anchor, epoch = setup()
    receive(event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC)), epoch)
    snapshot = load_report_snapshot(1, ('2026-09-01', '2026-09-30'), cache=None, subject_cache=None,
        repository=DeviceConfigRepository(), cancelled=Event())
    assert snapshot.data['subjects'][0].period_usage == 301
    assert sum(r.delta_value for r in snapshot.data['raw_counter_records'] if r.canonical_field == 'counter.user.total.total') == 301


def test_two_databases_converge_on_received_journal_without_financial_writes(tmp_path):
    anchor, epoch = setup()
    other = tmp_path / 'second.sqlite3'
    with sqlite3.connect(database.get_db_path()) as conn, sqlite3.connect(other) as target:
        conn.backup(target)
    events = [event_at(anchor, epoch, v, datetime(2026, 9, d, 9, tzinfo=UTC)) for v, d in [(200, 19), (401, 21)]]
    for event in events:
        receive(event, epoch)
    expected = load_subject_history([0])
    for event in reversed(events):
        receive(event, epoch, other)
    with database.report_read_database(other):
        assert load_subject_history([0]) == expected
        with database._open(commit=False) as conn:
            assert conn.execute('SELECT COUNT(*) FROM billing_ledger').fetchone()[0] == 0


@pytest.mark.parametrize('start,end', [(date(2026, 2, 1), date(2026, 2, 28)),
                                      (date(2026, 3, 1), date(2026, 3, 2)),
                                      (date(2026, 3, 2), date(2026, 3, 2))])
def test_month_context_matches_full_projection_without_parsing_unrelated_daily_json(start, end):
    anchor, epoch = setup(datetime(2026, 1, 1, 10, tzinfo=UTC))
    for at, value in [(datetime(2026, 1, 28, 10, tzinfo=UTC), 200),
                      (datetime(2026, 3, 1, 10, tzinfo=UTC), 400),
                      (datetime(2026, 3, 2, 10, tzinfo=UTC), 500)]:
        receive(event_at(anchor, epoch, value, at), epoch)
    expected = [r for r in shared() if start <= r.day <= end]
    graph = [r for r in shared(graph=True) if start <= r.day <= end]
    with database._open() as conn:
        conn.execute("UPDATE shared_report_days SET document_json='unread old detail',device_json='unread old detail' WHERE record_date<'2026-02-01'")
    assert list(shared(start=start, end=end)) == expected
    assert list(shared(start=start, end=end, graph=True)) == graph


def test_summary_failure_rolls_back_received_evidence_before_receipt():
    anchor, epoch = setup()
    with database._open() as conn:
        conn.execute("CREATE TRIGGER refuse_summary BEFORE UPDATE ON shared_report_months BEGIN SELECT RAISE(ABORT,'disk fixture'); END")
    event = event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC))
    with pytest.raises(sqlite3.DatabaseError, match='disk fixture'):
        receive(event, epoch)
    with database._open(commit=False) as conn:
        assert not conn.execute('SELECT 1 FROM web_observations WHERE event_id=?', (event['id'],)).fetchone()
        assert conn.execute('SELECT COUNT(*) FROM shared_report_events').fetchone()[0] == 1
    assert sum(r.delta for r in shared()) == 0


def test_used_replacement_preserves_each_epoch_delta_without_adding_its_opening_meter():
    anchor, epoch = setup(datetime(2026, 9, 18, 10, tzinfo=UTC))
    receive(event_at(anchor, epoch, 120, datetime(2026, 9, 18, 11, tzinfo=UTC)), epoch)
    DeviceConfigRepository().replace_all([DeviceConfig(0, 'Used replacement', '192.0.2.10', serial='USED')])
    at = datetime(2026, 9, 19, 10, tzinfo=UTC)
    store = WebObservationStore()
    event = store.record(0, reading(50000, at), serial='USED', printer_name='Used replacement', at=at)
    epochs = SlotEpochStore()
    replacement = epochs.confirm(epochs.preview(0, event['id'], operation='replacement', actor='A', reason='Reviewed used printer'))
    later = at + timedelta(hours=1)
    store.record(0, reading(50050, later), serial='USED', printer_name='Used replacement', at=later, epoch_id=replacement)
    rows = shared()
    assert sum(r.delta for r in rows) == 70
    assert sum(r.delta for r in rows if r.serial == 'USED') == 50
    subjects = load_subject_history([0], today=date(2026, 9, 20))
    assert sum(r.period_usage for r in subjects) == 70


def test_projection_cancellation_does_not_leave_request_cache_or_change_rows():
    from concurrent.futures import CancelledError
    anchor, epoch = setup()
    cancelled = Event()
    with database.report_read_database(database.get_db_path(), cancelled=cancelled):
        cancelled.set()
        with pytest.raises(CancelledError):
            shared()
    assert database.report_read_components() is None
    assert sum(r.delta for r in shared()) == 0


def test_ui_keeps_recorded_table_total_and_fractional_estimated_chart_separate():
    from PyQt6.QtCore import Qt
    from PyQt6.QtWidgets import QApplication
    from services.monthly_history import MonthlyUsage
    from ui.cloud_reports import CloudReportsWidget, USAGE_METRICS
    app = QApplication.instance() or QApplication([])
    note = 'ยอดสะสมช่วงขาดการอ่าน: fixture'
    def usage(day, amount, estimated=False):
        return MonthlyUsage(0, day, day, amount, 0, 0, 0, amount, 0, 0, 0, 1, 0,
            print_a4_bw=amount, print_available=True,
            remarks=(note, 'ค่าเฉลี่ยประมาณการเฉพาะวันที่ไม่มีข้อมูล Web') if estimated else (note,))
    sunday = usage(date(2026, 9, 20), 301)
    widget = CloudReportsWidget()
    widget.usage_granularity = 'D'
    widget.set_data(records=(sunday,), daily_records=(sunday,), usage_daily_records=(sunday,),
        chart_daily_records=(usage(date(2026, 9, 19), 150.5, True), usage(date(2026, 9, 20), 150.5, True)),
        monthly_records=(), subjects=(), devices=(DeviceConfig(0, 'Fixture', '192.0.2.10'),),
        period=('2026-09-01', '2026-09-30'), last_sync_at=None)
    widget.tabs.setCurrentIndex(3)
    widget._render_usage()
    app.processEvents()
    total_column = 1 + next(i for i, item in enumerate(USAGE_METRICS) if item[0] == 'total')
    assert widget.usage_monthly_table.rowCount() == 1
    assert widget.usage_monthly_table.item(0, total_column).text() == '301'
    chart = widget.usage_timeline_chart.chart()
    assert [p.y() for p in chart.series()[0].points()] == [150.5, 150.5]
    assert chart.axes(Qt.Orientation.Vertical)[0].labelFormat() == '%.2f'
    assert 'ประมาณการ' in chart.title()
    assert 'ตารางลงยอดวันก่อน' in widget.report_note_label.text()
    widget.close()


def test_received_history_exports_actual_allocation_and_missing_read_remark(tmp_path):
    from openpyxl import load_workbook
    from services.report_excel_export import CloudReportsExportData, export_cloud_reports_xlsx
    anchor, epoch = setup()
    receive(event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC)), epoch)
    raw = load_raw_counter_history([0])
    usage = load_usage_history([0], granularity='daily', today=date(2026, 9, 22))
    subjects = load_subject_history([0], today=date(2026, 9, 22))
    report = CloudReportsExportData(usage, usage, usage, (), subjects,
        tuple(DeviceConfigRepository().list()), ('2026-09-01', '2026-09-30'), None,
        datetime(2026, 9, 22, tzinfo=UTC), raw_counter_records=raw)
    path = export_cloud_reports_xlsx(tmp_path / 'shared-report.xlsx', report)
    workbook = load_workbook(path, data_only=True)
    try:
        notes = list(workbook['หมายเหตุรายละเอียด'].values)
        assert any(row[2] == '2026-09-20' and row[5] == 'ไม่มีจุดอ่านในวันนี้' and 'ยอดสะสมช่วงขาดการอ่าน' in row[6] for row in notes[1:])
        assert 'หมายเหตุรายงาน' in workbook.sheetnames
        machine = workbook['การใช้งานรายเครื่อง']
        assert machine.cell(5, machine.max_column - 1).value is None
        assert machine.cell(5, machine.max_column).value == 'ตรวจสอบ'
        daily = workbook['มุมมองรายวัน']
        assert daily.cell(5, daily.max_column).value is None
        rows = list(workbook['Legacy-รายวัน'].values)
        sunday = next(row for row in rows if isinstance(row[0], datetime) and row[0].date() == date(2026, 9, 20))
        assert 301 in sunday and 150.5 not in sunday
    finally:
        workbook.close()
    from services.report_models import build_printer_rows
    models = build_printer_rows(usage, tuple(DeviceConfigRepository().list()))
    assert models and all(not r['cost_available'] and r['review'] for r in models)


def test_archived_index_is_explicit_cancellable_and_resumes_without_duplicate_usage(monkeypatch):
    from services.web_observation_archive import WebObservationArchive
    from services import shared_report_index
    anchor, epoch = setup()
    receive(event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC)), epoch)
    while WebObservationArchive().archive_batch(today=date(2028, 1, 1)):
        pass
    with database._open() as conn:
        for table in ('shared_report_days', 'shared_report_months', 'shared_report_events'):
            conn.execute(f'DELETE FROM {table}')
    assert not shared()
    assert backfill_batch(database.get_db_path()) == 0
    cancelled = Event()
    original = shared_report_index.backfill_batch
    def first_only(path, **kwargs):
        value = original(path, limit=1, **kwargs)
        cancelled.set()
        return value
    monkeypatch.setattr(shared_report_index, 'backfill_batch', first_only)
    assert shared_report_index.index_archived_reports(database.get_db_path(), cancelled) == {'indexed': 1, 'cancelled': True}
    monkeypatch.setattr(shared_report_index, 'backfill_batch', original)
    cancelled.clear()
    assert shared_report_index.index_archived_reports(database.get_db_path(), cancelled)['indexed'] == 1
    assert shared_report_index.index_archived_reports(database.get_db_path(), cancelled)['indexed'] == 0
    assert sum(r.delta for r in shared()) == 301
    assert load_subject_history([0])[0].period_usage == 301


def test_known_but_unindexed_web_days_are_not_falsely_marked_as_days_without_readings():
    from services.shared_report_index import index_event, index_archived_reports
    from services.web_observation_archive import WebObservationArchive
    anchor, epoch = setup()
    saturday = event_at(anchor, epoch, 200, datetime(2026, 9, 19, 9, tzinfo=UTC))
    monday = event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC))
    receive(saturday, epoch)
    receive(monday, epoch)
    while WebObservationArchive().archive_batch(today=date(2028, 1, 1)):
        pass
    with database._open() as conn:
        for table in ('shared_report_days', 'shared_report_months', 'shared_report_events'):
            conn.execute(f'DELETE FROM {table}')
        reviewed = SlotEpochStore._epoch(conn, epoch)
        index_event(conn, reviewed, anchor)
        index_event(conn, reviewed, monday)
    partial = shared(graph=True)
    assert any('shared_report_index_incomplete' in r.issues for r in partial)
    assert not any(r.day == date(2026, 9, 19) and r.estimated for r in partial)
    index_archived_reports(database.get_db_path(), Event())
    complete = shared(graph=True)
    assert not any('shared_report_index_incomplete' in r.issues for r in complete)
    assert sum(r.delta for r in complete) == 301


def test_archive_panel_fills_ordinary_reports_only_after_explicit_button(monkeypatch):
    from services.web_observation_archive import WebObservationArchive
    from ui.archive_settings import ArchiveSettingsPanel
    from test_archive_settings_ui import ready
    anchor, epoch = setup()
    receive(event_at(anchor, epoch, 401, datetime(2026, 9, 21, 9, tzinfo=UTC)), epoch)
    while WebObservationArchive().archive_batch(today=date(2028, 1, 1)):
        pass
    with database._open() as conn:
        for table in ('shared_report_days', 'shared_report_months', 'shared_report_events'):
            conn.execute(f'DELETE FROM {table}')
    panel = ArchiveSettingsPanel(today=date(2028, 1, 1))
    ready(panel)
    assert panel.index_button.isEnabled() and panel.future is None
    assert not shared()
    panel.index_button.click()
    future = panel.future
    assert future is not None
    panel.index_button.click()
    assert panel.future is future
    future.result(timeout=10)
    panel._collect()
    ready(panel)
    assert '2' in panel.message.text()
    assert load_subject_history([0])[0].period_usage == 301
    with database._open(commit=False) as conn:
        assert conn.execute('SELECT COUNT(*) FROM billing_ledger').fetchone()[0] == 0
    panel.pool.shutdown(wait=True)
    panel.close()
