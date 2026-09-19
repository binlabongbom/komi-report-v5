from concurrent.futures import CancelledError
from datetime import datetime, timedelta, timezone
from fractions import Fraction
from functools import partial
from threading import Event
from time import monotonic

import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import QTimer
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication, QWidget

from services.web_usage_projection import CounterPoint, UsageInterval, CounterProjection
from services.web_evidence_export import export_web_evidence
from ui.report_export import ReportExportController


def projection(count=128):
    start = datetime(2026, 1, 1, tzinfo=timezone.utc)
    points = tuple(CounterPoint(start + timedelta(minutes=45*i), Fraction(1000+i), (f"fixture-{i}",)) for i in range(count))
    intervals = tuple(UsageInterval(a, b, Fraction(1), b.at.date(), (), ()) for a, b in zip(points, points[1:]))
    return CounterProjection("fixture-epoch", "fixture-revision", points, intervals, ())


@pytest.mark.parametrize("kind", ["xlsx", "pdf"])
@pytest.mark.parametrize("action", ["cancel", "locked", "fail"])
def test_incomplete_web_export_preserves_old_destination(tmp_path, monkeypatch, kind, action):
    destination = tmp_path / ("existing." + kind)
    destination.write_bytes(b"original")
    event = Event()
    if action == "locked":
        def locked(*args):
            raise PermissionError("File is open")
        monkeypatch.setattr("services.web_evidence_export.os.replace", locked)
    def progress(percent, _text):
        if percent >= 20 and action != "locked":
            if action == "fail":
                raise OSError("Synthetic failure")
            event.set()
    with pytest.raises((OSError, CancelledError)):
        export_web_evidence(destination, projection(), kind=kind, printer="Fixture", serial="fixture-serial",
                            counter_label="Total", cancelled=event, progress=progress)
    assert destination.read_bytes() == b"original"
    assert not list(tmp_path.glob(".komi-evidence-*"))


@pytest.mark.parametrize("kind", ["xlsx", "pdf"])
@pytest.mark.parametrize("count", [960, 4096])
def test_web_export_worker_keeps_ui_responsive(tmp_path, kind, count):
    app = QApplication.instance() or QApplication([])
    parent = QWidget()
    parent.show()
    controller = ReportExportController(parent)
    results, failures, ticks, stages = [], [], [], []
    controller.succeeded.connect(results.append)
    controller.failed.connect(failures.append)
    timer = QTimer(parent)
    timer.setInterval(10)
    timer.timeout.connect(lambda: ticks.append(monotonic()))
    timer.start()
    destination = tmp_path / ("evidence." + kind)
    def job(*, cancelled, progress):
        def update(percent, text):
            stages.append(percent)
            progress(percent, text)
        return export_web_evidence(destination, projection(count), kind=kind, printer="เครื่องทดสอบ <script>",
                                   serial="FIXTURE", counter_label="Total", cancelled=cancelled, progress=update)
    try:
        assert controller.start(job, title="Web evidence " + kind)
        assert not controller.start(job)
        deadline = monotonic() + 90
        while controller.busy and monotonic() < deadline:
            app.processEvents()
            QTest.qWait(10)
        assert not controller.busy and not failures and results == [destination]
        assert len(ticks) > 5
        assert max(b-a for a, b in zip(ticks, ticks[1:])) < 2
        assert stages == sorted(stages)
        if kind == "xlsx":
            book = load_workbook(destination, read_only=True, data_only=True)
            assert book["อ่านจริง"].max_row == count + 1
            assert sum(row[2] for row in book["ยอดตามช่วง"].iter_rows(min_row=2, values_only=True)) == count - 1
            book.close()
        else:
            assert destination.read_bytes().startswith(b"%PDF-")
    finally:
        timer.stop()
        controller.shutdown()
        parent.close()


@pytest.mark.parametrize("kind", ["xlsx", "pdf"])
@pytest.mark.parametrize("stage", [1, 95])
def test_cancel_before_write_and_before_publish_keeps_destination(tmp_path, kind, stage):
    path = tmp_path / ("keep." + kind)
    path.write_bytes(b"original")
    cancelled = Event()
    def progress(percent, _):
        if percent >= stage:
            cancelled.set()
    with pytest.raises(CancelledError):
        export_web_evidence(path, projection(), kind=kind, printer="Fixture", serial="Fixture",
                            counter_label="Total", cancelled=cancelled, progress=progress)
    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob(".komi-evidence-*"))


@pytest.mark.parametrize("error", [PermissionError("Folder is read-only"), OSError(28, "No space left")])
def test_destination_creation_error_is_reported_without_changing_existing_file(tmp_path, monkeypatch, error):
    path = tmp_path / "keep.xlsx"
    path.write_bytes(b"original")
    def fail(**_):
        raise error
    monkeypatch.setattr("services.web_evidence_export.tempfile.mkstemp", fail)
    with pytest.raises(OSError):
        export_web_evidence(path, projection(), kind="xlsx", printer="Fixture", serial="Fixture",
                            counter_label="Total", cancelled=Event(), progress=lambda *_: None)
    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob(".komi-evidence-*"))


def test_closing_evidence_window_cancels_its_worker_without_waiting():
    from ui.web_evidence_report import WebEvidenceDialog
    dialog = WebEvidenceDialog([])
    dialog.show()
    controller = dialog._ensure_export_controller()
    started, release = Event(), Event()
    def job(*, cancelled, progress):
        started.set()
        release.wait(5)
        if cancelled.is_set():
            raise CancelledError()
    try:
        assert controller.start(job)
        assert started.wait(2)
        before = monotonic()
        dialog.close()
        assert monotonic() - before < 1
        assert controller.closed and controller.cancelled.is_set()
        assert not controller.start(job)
        release.set()
        with pytest.raises(CancelledError):
            controller.future.result(timeout=3)
    finally:
        release.set()
        controller.shutdown()
        dialog.pool.shutdown(wait=True)
        dialog.close()
