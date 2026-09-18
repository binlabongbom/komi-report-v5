from datetime import date
from threading import Event

from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication
import pytest

import database
from services.device_config import DeviceConfigRepository
from services.web_evidence_report import load_evidence, inspect_counter, export_evidence_xlsx
from services.web_usage_projection import CounterIdentity
from services.web_observation_archive import WebObservationArchive
from test_web_usage_projection import events_for, BANGKOK, project
from ui.web_evidence_report import WebEvidenceDialog, projection_pdf


def request():
    return dict(slot=0, serial="SER-A", start=date(2026, 9, 18), end=date(2026, 9, 21), calendar_timezone=BANGKOK)


def test_range_filter_uses_instants_and_archived_payload_is_loaded_only_on_request(monkeypatch):
    events = events_for(("2026-09-17T23:00", 999), ("2026-09-18T17:00", 1000), ("2026-09-21T08:00", 1301), ("2026-09-22T00:00", 1400))
    archive = WebObservationArchive()
    while archive.archive_batch(today=date(2028, 1, 1)):
        pass
    result = load_evidence(database.get_db_path(), **request(), cancelled=Event())
    assert [e["id"] for e in result["events"]] == [e["id"] for e in events[1:3]]
    counter = CounterIdentity("counter.user.total.total", "user", "u-1")
    projection = inspect_counter(result, counter, **request())
    assert projection.total_usage == 301
    assert projection.intervals[0].estimated_per_day == 150.5


def test_bounded_explicit_request_does_not_silently_truncate(monkeypatch):
    events_for(("2026-09-18T17:00", 1000), ("2026-09-21T08:00", 1301))
    monkeypatch.setattr("services.web_evidence_report.MAX_REPORT_EVENTS", 1)
    with pytest.raises(ValueError, match="ช่วงวันที่สั้นลง"):
        load_evidence(database.get_db_path(), **request(), cancelled=Event())


def test_exports_preserve_decimal_estimates_remarks_and_escape_labels(tmp_path):
    result = project(events_for(("2026-09-18T17:00", 1000), ("2026-09-21T08:00", 1301)))
    xlsx = tmp_path / "evidence.xlsx"
    export_evidence_xlsx(xlsx, result, printer="=1+1", serial="SER-A", counter_label="@UNTRUSTED")
    workbook = load_workbook(xlsx)
    assert workbook["ขอบเขตข้อมูล"]["B1"].data_type == "s"
    assert workbook["ขอบเขตข้อมูล"]["B1"].value == "'=1+1"
    graph = workbook["กราฟรายวัน"]
    assert graph["D2"].value == graph["D3"].value == 150.5
    assert graph["E2"].value == "ประมาณการ"
    assert "ไม่ใช่ยอดที่ยืนยัน" in workbook["ยอดตามช่วง"]["G2"].value
    workbook.close()
    pdf = tmp_path / "evidence.pdf"
    projection_pdf(pdf, result, title="Fixture <script> · รายงานหลักฐาน")
    assert pdf.read_bytes().startswith(b"%PDF-")
    assert pdf.stat().st_size > 2000


def test_dialog_loads_only_on_click_and_resets_exports_after_date_change():
    events_for(("2026-09-18T17:00", 1000), ("2026-09-21T08:00", 1301))
    dialog = WebEvidenceDialog(DeviceConfigRepository().list(enabled_only=False), today=date(2026, 9, 21), calendar_timezone=BANGKOK)
    assert dialog.future is None and dialog.evidence is None
    assert not dialog.xlsx_button.isEnabled()
    dialog.load()
    dialog.future.result(timeout=10)
    dialog.collect()
    dialog.future.result(timeout=10)
    dialog.collect()
    assert dialog.projection is not None
    assert dialog.xlsx_button.isEnabled() and dialog.pdf_button.isEnabled()
    assert dialog.table.rowCount() == 1
    assert len(dialog.cumulative.chart().series()) == 2
    dialog.start_date.setDate(dialog.start_date.date().addDays(1))
    assert dialog.projection is None
    assert not dialog.xlsx_button.isEnabled()
    dialog.pool.shutdown(wait=True)


def test_missing_history_shows_unavailable_instead_of_zero_result():
    events_for(("2026-08-01T17:00", 1000))
    dialog = WebEvidenceDialog(DeviceConfigRepository().list(enabled_only=False), today=date(2026, 9, 21), calendar_timezone=BANGKOK)
    dialog.load()
    dialog.future.result(timeout=10)
    dialog.collect()
    assert dialog.projection is None
    assert "ไม่มีผล Web" in dialog.status.text()
    assert not dialog.xlsx_button.isEnabled()
    dialog.pool.shutdown(wait=True)


def test_reviewed_reset_epochs_are_separate_in_projection_dialog_and_export(tmp_path):
    from datetime import datetime, UTC
    from services.web_observations import WebObservationStore
    from test_slot_epoch_store import fixture, confirm
    from test_subject_history import _reading
    store, _, initial = fixture()
    at = datetime(2026, 9, 18, 10, tzinfo=UTC)
    events = WebObservationStore()
    boundary = events.record(0, _reading(50, collected_at=at), at=at,
        serial="OLD", printer_name="Fixture", epoch_id=initial)
    reset = confirm(store, boundary, "reset")
    at = datetime(2026, 9, 19, 10, tzinfo=UTC)
    events.record(0, _reading(80, collected_at=at), at=at,
        serial="OLD", printer_name="Fixture", epoch_id=reset)
    args = dict(slot=0, serial="OLD", start=date(2026, 9, 16), end=date(2026, 9, 19), calendar_timezone=BANGKOK)
    evidence = load_evidence(database.get_db_path(), **args, cancelled=Event())
    counter = CounterIdentity("counter.user.total.total", "user", "u-1")
    assert set(evidence["epochs"]) == {initial, reset}
    with pytest.raises(ValueError, match="เลือกช่วงฐาน"):
        inspect_counter(evidence, counter, **args)
    projection = inspect_counter(evidence, counter, **args, epoch_id=reset)
    assert projection.total_usage == 30
    assert [point.value for point in projection.points] == [50, 80]
    exported = tmp_path / "reset.xlsx"
    export_evidence_xlsx(exported, projection, printer="Fixture", serial="OLD", counter_label="User")
    workbook = load_workbook(exported)
    assert dict(workbook["ขอบเขตข้อมูล"].values)["ช่วงฐานมิเตอร์"] == reset
    workbook.close()
    dialog = WebEvidenceDialog(DeviceConfigRepository().list(enabled_only=False), today=date(2026, 9, 19), calendar_timezone=BANGKOK)
    dialog.load()
    dialog.future.result(timeout=10)
    dialog.collect()
    dialog.future.result(timeout=10)
    dialog.collect()
    assert dialog.epochs.count() == 2 and dialog.epochs.currentData() == reset
    assert dialog.projection.total_usage == 30
    dialog.epochs.setCurrentIndex(dialog.epochs.findData(initial))
    assert not dialog.xlsx_button.isEnabled()
    dialog.future.result(timeout=10)
    dialog.collect()
    assert dialog.projection.epoch_id == initial
    assert "lower_counter_or_clock_conflict" in dialog.projection.issues
    dialog.pool.shutdown(wait=True)


def test_projector_rejects_bound_events_from_other_epoch():
    from datetime import datetime, UTC
    from services.web_observations import WebObservationStore
    from services.web_usage_projection import project_counter
    from test_slot_epoch_store import fixture
    from test_subject_history import _reading
    _, _, epoch = fixture()
    at = datetime(2026, 9, 18, 10, tzinfo=UTC)
    event = WebObservationStore().record(0, _reading(2010, collected_at=at), at=at,
        serial="OLD", printer_name="Fixture", epoch_id=epoch)
    with pytest.raises(ValueError, match="different reviewed"):
        project_counter([event], slot=0, serial="OLD", epoch_id="unrelated",
            counter=CounterIdentity("counter.user.total.total", "user", "u-1"), calendar_timezone=BANGKOK)
