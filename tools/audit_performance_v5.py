"""Source-only Qt performance audit on verified synthetic data, never customer data.

Run each mode in a fresh process. No services, network, release build or external
process may be launched by the measured process. Results are flushed after each
step so an interrupted expensive operation does not erase earlier measurements.
"""
from __future__ import annotations

import argparse
from contextlib import closing
import faulthandler
from datetime import datetime, timedelta, timezone
from fractions import Fraction
import json
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'source'), str(ROOT / 'source/tests'), str(ROOT / 'tools')]
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
os.environ['KOMI_DISCOVERY_DISABLED'] = '1'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--mode', choices=('ui', 'queries', 'evidence', 'charts', 'export'), required=True)
    parser.add_argument('--synthetic-db', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--days', type=int, choices=(30, 128), default=30)
    parser.add_argument('--format', choices=('xlsx', 'pdf'), default='xlsx')
    parser.add_argument('--hold-previous-chart', action='store_true',
                        help='Diagnostic control only: retain the old Python chart wrapper through setChart')
    parser.add_argument('--report-years', type=int, choices=(1, 5), default=1)
    args = parser.parse_args()
    faulthandler.enable()
    faulthandler.dump_traceback_later(60, repeat=True)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    result = {'synthetic': True, 'mode': args.mode, 'started_at': datetime.now().isoformat(),
              'python': sys.version, 'platform': sys.platform, 'measurements': [],
              'blocked_operations': [], 'complete': False,
              'scope': 'Source/offscreen; no Host, real network, Web service, installer or release build'}

    def save():
        args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')

    def no_external(event, _values):
        if event in ('socket.connect', 'socket.sendto', 'socket.bind', 'subprocess.Popen', 'os.system'):
            result['blocked_operations'].append(event)
            if len(result['blocked_operations']) <= 6:
                import traceback
                result.setdefault('blocked_operation_sites', []).append({
                    'event': event,
                    'stack': [{'file': frame.filename, 'line': frame.lineno, 'function': frame.name}
                              for frame in traceback.extract_stack(limit=6)[:-1]]})
            raise PermissionError('Performance audit forbids external operations')

    sys.addaudithook(no_external)
    from measure_idle_runtime import counters
    from PyQt6.QtCore import PYQT_VERSION_STR, QT_VERSION_STR, QTimer
    from PyQt6.QtWidgets import QApplication
    from PyQt6.QtTest import QTest
    from qt_render_environment import configure_offscreen_fonts
    app = QApplication([])
    app.setQuitOnLastWindowClosed(False)
    result['qt_version'] = QT_VERSION_STR
    result['pyqt_version'] = PYQT_VERSION_STR
    configure_offscreen_fonts(app)
    beats = []
    timer = QTimer()
    timer.setInterval(10)
    timer.timeout.connect(lambda: beats.append(time.perf_counter()))
    timer.start()

    def step(name, action, ready=None, timeout=180):
        QTest.qWait(30)
        beats.clear()
        before = counters()
        started = time.perf_counter()
        result['in_progress'] = name
        save()
        value = action()
        action_seconds = time.perf_counter() - started
        while ready is not None and not ready():
            if time.perf_counter() - started > timeout:
                raise TimeoutError(f'{name}: not ready after {timeout} seconds')
            QTest.qWait(5)
        finished = time.perf_counter()
        # Include the first post-operation timeout to expose synchronous stalls.
        QTest.qWait(25)
        after = counters()
        samples = [started] + beats
        gaps = [b - a for a, b in zip(samples, samples[1:])]
        row = {'name': name, 'seconds_until_ready': finished - started,
               'action_seconds': action_seconds, 'heartbeat_count': len(beats),
               'max_event_loop_gap_ms': max(gaps, default=0) * 1000,
               'gaps_above_100ms': sum(g > .1 for g in gaps),
               'before': before, 'after': after}
        if isinstance(value, dict):
            row['details'] = value
        result['measurements'].append(row)
        result['in_progress'] = None
        save()
        print(json.dumps({k: row[k] for k in ('name', 'seconds_until_ready', 'max_event_loop_gap_ms')}), flush=True)
        return value

    temporary = tempfile.TemporaryDirectory(prefix='komi-performance-audit-')
    directory = Path(temporary.name)
    import database
    database.set_db_path(directory / 'synthetic.sqlite3')
    if args.mode in ('ui', 'queries', 'export'):
        if args.synthetic_db is None:
            parser.error('--synthetic-db is required for ui/queries')
        source = args.synthetic_db.resolve(strict=True)
        with closing(sqlite3.connect(source.as_uri() + '?mode=ro', uri=True)) as original:
            devices = original.execute('SELECT slot,serial,host FROM devices ORDER BY slot').fetchall()
            expected = [(slot, f'FIXTURE-{slot}', f'192.0.2.{slot + 10}') for slot in range(3)]
            if devices != expected:
                raise ValueError('Only the documented three-printer five-year fixture is accepted')
            result['fixture'] = {'source': str(source), 'original_opened_read_only': True,
                                 'devices': devices, 'database_bytes': source.stat().st_size,
                                 'users': original.execute('SELECT COUNT(DISTINCT subject_id) FROM daily_subject_records').fetchone()[0],
                                 'date_range': original.execute('SELECT MIN(record_date),MAX(record_date) FROM daily_subject_records').fetchone(),
                                 'row_counts': {name: original.execute(f'SELECT COUNT(*) FROM {name}').fetchone()[0]
                                    for name in ('daily_meter_records', 'daily_subject_records', 'daily_raw_counter_records')}}
            with closing(sqlite3.connect(database.get_db_path())) as target:
                original.backup(target)
    initialized = time.perf_counter()
    database.init_db()
    result['copy_migration_init_seconds'] = time.perf_counter() - initialized
    database.set_cfg('web_service_enabled', '0')
    database.set_cfg('ui_language', 'th')
    with closing(sqlite3.connect(database.get_db_path())) as connection:
        result['copy_schema'] = connection.execute('PRAGMA user_version').fetchone()[0]
    save()

    if args.mode == 'ui':
        from ui_consistency import install_button_system
        from app import MainWindow
        from services.report_snapshot import SnapshotDataCache
        original_read = SnapshotDataCache.read
        result['snapshot_reads'] = []
        def timed_read(self, key, loader, cancelled):
            def load():
                started = time.perf_counter()
                value = loader()
                result['snapshot_reads'].append({'component': str(key[1:]),
                                                'seconds': time.perf_counter()-started})
                return value
            return original_read(self, key, load, cancelled)
        SnapshotDataCache.read = timed_read
        install_button_system(app)
        # Offscreen exposes an artificial 800x800 screen; use a declared desktop
        # viewport without changing widget/layout constraints or loading code.
        MainWindow._fit_to_available_screen = lambda self, **kw: self.resize(1440, 900)
        result['viewport'] = [1440, 900]
        holder = {}

        def construct():
            holder['window'] = MainWindow()
            holder['window'].show()
        step('main_window_construct_show', construct)
        window = holder['window']
        window._startup_refresh_pending = False
        step('initial_refresh', window._refresh_all,
             lambda: not window.pg_overview.summary_reader.busy)

        def reports_ready():
            page = window.pg_history
            if not hasattr(page, '_snapshot_jobs'):
                return False
            center = page.report_center
            return (not page._snapshot_jobs and not page._report_refresh_timer.isActive()
                    and not center._loading and not center._pending_render_tail
                    and not center._render_tail_timer.isActive())

        def navigate(index):
            window._nav(window.nav_btns[index])

        def page_ready(index):
            return index not in window._deferred_pages

        for index, name in ((1, 'devices'), (2, 'reports'), (3, 'settings')):
            step(f'first_navigation_{name}', lambda i=index: navigate(i),
                 reports_ready if index == 2 else lambda i=index: page_ready(i))
        navigate(2)
        QTest.qWait(100)
        page = window.pg_history
        center = page.report_center
        result['report_tabs'] = [center.tabs.tabText(i) for i in range(center.tabs.count())]

        def report_rows():
            return {'period': center.current_period(), 'cache_entries': len(page._snapshot_cache),
                    'selected_tab': center.tabs.currentIndex(),
                    'cached_snapshots': [{
                        'key': str(key),
                        'rows': {k: len(v) for k, v in snapshot.data.items() if isinstance(v, (list, tuple))}
                    } for key, snapshot in page._snapshot_cache.items()]}

        for index in range(center.tabs.count()):
            step(f'report_month_tab_{index}', lambda i=index: center.tabs.setCurrentIndex(i), reports_ready)
        period = ('2021-09-18' if args.report_years == 5 else '2025-09-18', '2026-09-17')
        result['selected_history_years'] = args.report_years

        def year_users():
            center.tabs.setCurrentIndex(1)
            center.set_periods((), period)
            page._staged_period_changed(period)
        step('report_year_users', year_users, reports_ready)
        result['year_users_data'] = report_rows()
        snapshot = page._snapshot_cache[(1, period)]
        with database._open(commit=False) as connection:
            expected = dict(connection.execute('''SELECT subject_name,SUM(delta_total)
                FROM daily_subject_records WHERE record_date BETWEEN ? AND ? GROUP BY subject_name''', period))
            expected_raw = connection.execute('''SELECT COUNT(*) FROM daily_raw_counter_records
                WHERE record_date BETWEEN ? AND ?''', period).fetchone()[0]
        users = snapshot.data['prepared_models']['users']
        actual = {row['name']: row['usage'] for row in users}
        assert actual == expected, 'Preview user totals do not match selected source dates'
        assert len(snapshot.data['raw_counter_records']) == expected_raw
        visible_error = center.error_label.text() if center.error_label.isVisible() else ''
        assert not visible_error, visible_error
        result['preview_validation'] = {'users': len(users), 'totals_match': True,
            'selected_raw_rows': expected_raw, 'period': period}
        save()
        for index in (0, 2, 3, 4):
            step(f'report_year_tab_{index}', lambda i=index: center.tabs.setCurrentIndex(i), reports_ready)
        result['year_all_tabs_data'] = report_rows()
        raw_snapshots = [snapshot.data['raw_counter_records'] for snapshot in page._snapshot_cache.values()
                         if snapshot.data.get('raw_counter_records')]
        result['raw_snapshot_objects'] = len({id(rows) for rows in raw_snapshots})
        for cycle in range(5):
            for index in range(center.tabs.count()):
                step(f'warm_cycle_{cycle}_tab_{index}', lambda i=index: center.tabs.setCurrentIndex(i), reports_ready)
        result['report_error'] = center.error_label.text() if center.error_label.isVisible() else ''
        from ui.archive_settings import ArchiveSettingsPanel
        archive_panel = ArchiveSettingsPanel()
        archive_panel.resize(700, 450)
        step('archive_metadata_first', archive_panel.show, lambda: not archive_panel.metadata_reader.busy)
        step('archive_metadata_warm', archive_panel.refresh, lambda: not archive_panel.metadata_reader.busy)
        result['archive_metadata'] = archive_panel.status.text()
        archive_panel.metadata_reader.shutdown()
        archive_panel.close()
        archive_panel.deleteLater()
        window._quitting = True
        window.close()
        page._snapshot_pool.shutdown(wait=True, cancel_futures=True)
        window.deleteLater()
        QTest.qWait(100)
    elif args.mode == 'export':
        from datetime import date
        from ui.report_export import ReportExportController
        from PyQt6.QtWidgets import QWidget, QPushButton
        from services.report_export_job import export_report_selection
        from services.device_config import DeviceConfigRepository
        from services.monthly_history import load_usage_history
        from openpyxl import load_workbook
        parent = QWidget()
        button = QPushButton('Navigation response', parent)
        parent.show()
        controller = ReportExportController(parent)
        errors, completed, phases, clicks = [], [], [], []
        controller.failed.connect(errors.append)
        controller.succeeded.connect(completed.append)
        button.clicked.connect(lambda: clicks.append(time.perf_counter()))
        timer.timeout.connect(button.click)
        path = args.output.with_suffix('.xlsx')
        period_start = date(2021 if args.report_years == 5 else 2025, 9, 18)
        period_end = date(2026, 9, 17)
        days = (period_end - period_start).days + 1
        with database._open(commit=False) as conn:
            fixture_users = conn.execute('SELECT COUNT(DISTINCT subject_id) FROM daily_subject_records').fetchone()[0]
        database_path = database.get_db_path()
        def job(*, cancelled, progress):
            def phase(percent, message):
                phases.append({'percent': percent, 'seconds': time.perf_counter()-begun, 'message': message})
                print(json.dumps({'export_progress': phases[-1]}), flush=True)
                progress(percent, message)
            # Include all selection preparation, DB/Archive reads and writing.
            from database import report_read_database
            with report_read_database(database_path, cancelled=cancelled):
                devices = DeviceConfigRepository().list(enabled_only=False)
                monthly = load_usage_history([0, 1, 2], granularity='monthly')
                return export_report_selection(path, period_start, period_end,
                    database_path=database_path, devices=devices, monthly_records=monthly,
                    cancelled=cancelled, progress=phase)
        begun = time.perf_counter()
        step(f'export_{args.report_years}_years_from_database', lambda: controller.start(job), lambda: not controller.busy,
             timeout=900)
        assert not errors, errors
        assert completed == [path]
        with path.open('rb') as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            rows = list(book['มุมมองผู้ใช้'].iter_rows(min_row=5, values_only=True))
            assert len(rows) == fixture_users
            totals = {row[0]: row[-2] for row in rows}
            expected = {f'Fixture user {i}': (10 + i % 5) * days * 3 for i in range(fixture_users)}
            assert totals == expected, (list(totals.items())[:2], list(expected.items())[:2])
            result['validation'] = {'users': len(rows), 'totals_match': True,
                'sheets': book.sheetnames, 'source_raw_rows': fixture_users * days * 3,
                'period': [str(period_start), str(period_end)],
                'archive_state': 'hot five-year fixture; archive path checked separately'}
            book.close()
        result['export_phases'] = phases
        result['navigation_clicks'] = len(clicks)
        result['output_bytes'] = path.stat().st_size
        controller.shutdown()
        parent.close()
        parent.deleteLater()
    elif args.mode == 'charts':
        import gc
        import weakref
        from PyQt6 import sip
        from PyQt6.QtCharts import QChartView
        from PyQt6.QtCore import QCoreApplication, QEvent
        from ui.cloud_reports import CloudReportsWidget
        from excel_export_fixture import report_for
        replace_chart = QChartView.setChart
        ownership = []
        def observe_replace(view, chart):
            previous = view.chart()
            replace_chart(view, chart)
            ownership.append({'previous_deleted': sip.isdeleted(previous),
                              'previous_python_owned': sip.ispyowned(previous),
                              'previous_has_parent': previous.parent() is not None})
        if args.hold_previous_chart:
            QChartView.setChart = observe_replace
        result['chart_replacement_control'] = args.hold_previous_chart
        report = report_for(2)
        cloud = CloudReportsWidget(staged=True)
        cloud.resize(1440, 900)
        cloud.show()
        alive = set()
        destroyed = []
        watched_count = 0
        chart_refs = []
        def watch_charts():
            nonlocal watched_count
            for view in cloud.findChildren(QChartView):
                chart = view.chart()
                address = int(sip.unwrapinstance(chart))
                if address in alive:
                    continue
                alive.add(address)
                watched_count += 1
                chart_refs.append(weakref.ref(chart))
                def released(*_args, token=address):
                    alive.discard(token)
                    destroyed.append(token)
                chart.destroyed.connect(released)
        def render():
            cloud.set_data(records=report.selected_records, daily_records=report.selected_daily_records,
                usage_daily_records=report.usage_daily_records, monthly_records=report.monthly_records,
                subjects=report.subjects, devices=report.devices, period=report.period,
                raw_counter_records=report.raw_counter_records, tab_index=0, last_sync_at=None)
        watch_charts()
        for cycle in range(10):
            step(f'replace_overview_charts_{cycle}', render,
                 lambda: not cloud._pending_render_tail and not cloud._render_tail_timer.isActive())
            watch_charts()
        gc.collect()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        result['charts_before_close'] = {'watched': watched_count, 'destroyed': len(destroyed),
                                        'alive': len(alive), 'views': len(cloud.findChildren(QChartView))}
        cloud.close()
        cloud.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        gc.collect()
        result['charts_after_close'] = {'watched': watched_count, 'destroyed': len(destroyed), 'alive': len(alive)}
        result['charts_after_close']['live_wrappers_not_deleted'] = sum(
            ref() is not None and not sip.isdeleted(ref()) for ref in chart_refs)
        result['replacement_ownership'] = ownership
        QChartView.setChart = replace_chart
        result['fixture'] = '90 synthetic users, three printers, two days, ten Overview data replacements; QObject destroyed signals; callbacks hold only integer pointers'
    elif args.mode == 'queries':
        from datetime import date
        from services.monthly_history import load_monthly_usage
        from services.subject_history import load_subject_history
        from services.raw_counter_history import load_raw_counter_history
        for name, action in (
            ('monthly_all_five_years', lambda: load_monthly_usage([0, 1, 2], today=date(2026, 9, 18))),
            ('subjects_30days', lambda: load_subject_history([0, 1, 2], period_start=date(2026, 8, 19), period_end=date(2026, 9, 17))),
            ('raw_30days', lambda: load_raw_counter_history([0, 1, 2], period_start=date(2026, 8, 19), period_end=date(2026, 9, 17))),
            ('subjects_365days', lambda: load_subject_history([0, 1, 2], period_start=date(2025, 9, 18), period_end=date(2026, 9, 17))),
            ('raw_365days', lambda: load_raw_counter_history([0, 1, 2], period_start=date(2025, 9, 18), period_end=date(2026, 9, 17))),
        ):
            step(name, lambda a=action: {'returned_rows': len(a())})
        from services.history_archive import HistoryArchive
        step('archive_settings_stats', lambda: HistoryArchive().stats(today=date(2026, 9, 19)))
        import services.summary_service as summary_service
        subject_loader = summary_service.load_subject_history
        summary_subject_times = []
        def timed_subject_loader(*a, **kw):
            started = time.perf_counter()
            rows = subject_loader(*a, **kw)
            summary_subject_times.append({'seconds': time.perf_counter() - started, 'rows': len(rows),
                                          'kwargs': {k: str(v) for k, v in kw.items()}})
            return rows
        summary_service.load_subject_history = timed_subject_loader
        step('overview_month_summary', lambda: {
            'top_users': len(summary_service.build_overview_summary('month', now=datetime(2026, 9, 19))['top_users'])})
        result['overview_subject_load'] = summary_subject_times
        result['query_note'] = 'Service calls measured synchronously only for timing; production staged Reports executes these reads in workers.'
    else:
        from services.web_usage_projection import CounterPoint, CounterProjection, UsageInterval
        from services.web_evidence_report import MAX_REPORT_EVENTS
        from services.web_evidence_export import export_web_evidence
        from ui.web_evidence_report import WebEvidenceDialog
        from ui.report_export import ReportExportController
        from PyQt6.QtWidgets import QWidget, QPushButton
        from functools import partial
        first = datetime(2025, 9, 18, 8, tzinfo=timezone(timedelta(hours=7)))
        points = []
        for day in range(args.days):
            for sample in range(32):
                at = first + timedelta(days=day, minutes=sample * 15)
                points.append(CounterPoint(at, Fraction(10000 + len(points) * 10), (f'fixture-{day}-{sample}',)))
        intervals = tuple(UsageInterval(a, b, b.value - a.value, b.at.date(), (), ())
                          for a, b in zip(points, points[1:]))
        projection = CounterProjection('fixture-epoch', 'fixture-revision', tuple(points), intervals, ())
        assert len(points) <= MAX_REPORT_EVENTS
        path = directory / f'evidence.{args.format}'
        result['fixture'] = {'days': args.days, 'reads_per_day': 32, 'points': len(points),
                             'intervals': len(intervals), 'counters': 1, 'format': args.format}
        parent = QWidget()
        button = QPushButton('Navigate', parent)
        parent.show()
        controller = ReportExportController(parent)
        errors, completed, clicks = [], [], []
        controller.failed.connect(errors.append)
        controller.succeeded.connect(completed.append)
        button.clicked.connect(lambda: clicks.append(True))
        timer.timeout.connect(button.click)
        step(f'evidence_{args.format}_{args.days}days', lambda: controller.start(partial(
            export_web_evidence, path, projection, kind=args.format,
            printer='Fixture', serial='FIXTURE-0', counter_label='Total')), lambda: not controller.busy)
        assert not errors, errors
        assert completed == [path]
        result['navigation_clicks'] = len(clicks)
        if args.format == 'xlsx':
            from openpyxl import load_workbook
            with path.open('rb') as stream:
                book = load_workbook(stream, read_only=True)
                result['validation'] = {sheet.title: sheet.max_row for sheet in book}
                assert book.worksheets[1].max_row == len(points) + 1
                assert book.worksheets[2].max_row == len(intervals) + 1
                book.close()
        else:
            assert path.read_bytes().startswith(b'%PDF-')
            result['validation'] = 'PDF header and non-empty output; visual pagination not audited here'
            import shutil
            pdf_copy = args.output.with_suffix('.pdf')
            shutil.copyfile(path, pdf_copy)
            result['pdf_for_visual_check'] = str(pdf_copy)
        controller.shutdown()
        parent.close()
        parent.deleteLater()
        result['output_bytes'] = path.stat().st_size
    result['complete'] = True
    save()
    faulthandler.cancel_dump_traceback_later()
    timer.stop()
    app.quit()
    # Qt reports workers are shut down by window.close; clear owned objects before
    # allowing TemporaryDirectory to remove the disposable migrated copy.
    temporary.cleanup()
    result['temporary_copy_cleaned'] = True
    save()


if __name__ == '__main__':
    main()
