"""Measure synthetic Excel exports, never build the app or contact a printer."""
import argparse
import json
from pathlib import Path
import sys
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'source'), str(ROOT / 'source/tests')]
from excel_export_fixture import report_for
from services.report_excel_export import export_cloud_reports_xlsx
from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument('--days', type=int, default=30)
parser.add_argument('--name', default='current')
parser.add_argument('--ui', action='store_true')
args = parser.parse_args()
folder = ROOT / 'outputs/excel-export-qa'
folder.mkdir(parents=True, exist_ok=True)
report = report_for(days=args.days)
started = perf_counter()
ui_metrics = {}
if args.ui:
    import os
    os.environ['QT_QPA_PLATFORM'] = 'offscreen'
    from functools import partial
    from PyQt6.QtCore import QTimer
    from PyQt6.QtTest import QTest
    from PyQt6.QtWidgets import QApplication, QPushButton, QWidget
    from ui.report_export import ReportExportController
    app = QApplication([])
    parent = QWidget()
    button = QPushButton('Navigate', parent)
    parent.show()
    controller = ReportExportController(parent)
    ticks, clicks, errors, results = [], [], [], []
    button.clicked.connect(lambda: clicks.append(1))
    timer = QTimer(parent)
    timer.setInterval(10)
    timer.timeout.connect(lambda: ticks.append(perf_counter()))
    timer.timeout.connect(button.click)
    controller.failed.connect(errors.append)
    controller.succeeded.connect(results.append)
    timer.start()
    controller.start(partial(export_cloud_reports_xlsx, folder / (args.name + '.xlsx'), report))
    while controller.busy:
        QTest.qWait(10)
    timer.stop()
    controller.shutdown()
    parent.close()
    if errors:
        raise errors[0]
    result = results[0]
    ui_metrics = {'gui_timer_ticks': len(ticks), 'navigation_clicks': len(clicks),
                  'maximum_tick_gap_ms': round(max(b-a for a,b in zip(ticks,ticks[1:]))*1000, 1)}
else:
    result = export_cloud_reports_xlsx(folder / (args.name + '.xlsx'), report)
elapsed = perf_counter() - started
workbook = load_workbook(result, read_only=True, data_only=True)
sheet = workbook['มุมมองผู้ใช้']
rows = list(sheet.iter_rows(min_row=5, values_only=True))
assert len(rows) == 90
assert all(row[-2] == args.days * 8 * 3 for row in rows)
assert all(row[27] == args.days * 16 for row in rows)
workbook.close()
summary = {'synthetic': True, 'users': 90, 'printers': 3, 'days': args.days,
           'raw_rows': len(report.raw_counter_records), 'seconds': round(elapsed, 3),
           'bytes': result.stat().st_size, 'verified_user_rows': len(rows), 'totals_verified': True}
summary.update(ui_metrics)
(folder / (args.name + '.json')).write_text(json.dumps(summary, indent=2), encoding='utf-8')
print(json.dumps(summary))
