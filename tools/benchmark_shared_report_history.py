"""Synthetic source-only report benchmark; never opens the installed customer DB.

Run with PYTHONPATH=source. No listener, installer, Host request or real printer.
One sample/day gives the same compact daily row count as frequent daily reads.
"""
from datetime import UTC, date, datetime, timedelta
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from time import perf_counter

import database
from komi_snmp_profiler.models import MeterReading, WebCounterValue, Confidence, WebAuthState
from services.device_config import DeviceConfigRepository
from services.raw_counter_history import load_raw_counter_history
from services.report_snapshot import load_report_snapshot
from services.slot_epoch_store import SlotEpochStore
from services.subject_history import load_subject_history
from services.web_observations import WebObservationStore


def measure_export(path, root, *, users, days, raw_rows):
    """Exercise the same progress controller and worker used by Reports."""
    from functools import partial
    import gc
    from PyQt6.QtCore import QTimer, Qt
    from PyQt6.QtTest import QTest
    from PyQt6.QtWidgets import QApplication, QPushButton
    from openpyxl import load_workbook
    from services.report_export_job import export_report_selection
    from ui.report_export import ReportExportController

    app = QApplication.instance() or QApplication([])
    button = QPushButton('Synthetic responsiveness probe')
    button.show()
    controller = ReportExportController(button)
    results, errors, times, progress, clicks, pauses, gaps = [], [], [], [], [], [], []
    gc_started = [0.0]
    def collection(phase, info):
        if phase == 'start':
            gc_started[0] = perf_counter()
        elif phase == 'stop':
            pauses.append(perf_counter() - gc_started[0])
    controller.succeeded.connect(results.append)
    controller.failed.connect(errors.append)
    button.clicked.connect(lambda: clicks.append(perf_counter()))
    timer = QTimer()
    timer.setInterval(20)
    def tick():
        now = perf_counter()
        previous_progress = progress[-1] if progress else 0
        gaps.append((now-times[-1], previous_progress, len(times)))
        times.append(now)
        if controller.dialog is not None:
            progress.append(controller.dialog.progress_bar.value())
        QTest.mouseClick(button, Qt.MouseButton.LeftButton)
    timer.timeout.connect(tick)
    target = Path(root) / 'synthetic-year.xlsx'
    start = date(2025, 1, 1)
    began = perf_counter()
    times.append(began)
    try:
        gc.callbacks.append(collection)
        controller.start(partial(export_report_selection, target, start, start+timedelta(days=days-1),
            database_path=path, devices=tuple(DeviceConfigRepository().list())))
        timer.start()
        deadline = began + 240
        while controller.busy and perf_counter() < deadline:
            QTest.qWait(5)
        elapsed = perf_counter() - began
        times.append(perf_counter())
        assert not controller.busy and results and not errors, errors or 'Export timed out'
        assert clicks and len(set(progress)) > 3 and max(progress) >= 90
    finally:
        gc.callbacks.remove(collection)
        timer.stop()
        controller.shutdown()
        controller.pool.shutdown(wait=True, cancel_futures=True)
        button.close()
        app.processEvents()
    workbook = load_workbook(target, read_only=True, data_only=True)
    try:
        names = [row[0] for row in workbook['มุมมองผู้ใช้'].iter_rows(min_row=5, max_col=1, values_only=True)]
        assert sum(isinstance(name, str) and name.startswith('Synthetic user ') for name in names) == users
        assert workbook['หมายเหตุรายละเอียด'].max_row == raw_rows + 1
    finally:
        workbook.close()
    return {'export_seconds': round(elapsed, 3), 'ui_probe_clicks': len(clicks),
            'ui_max_gap_ms': round(max(b-a for a, b in zip(times, times[1:])) * 1000, 1),
            'ui_largest_gaps': [{'ms': round(gap*1000, 1), 'progress_before': phase, 'sample': sample}
                                for gap, phase, sample in sorted(gaps, reverse=True)[:5]],
            'gc_max_pause_ms': round(max(pauses, default=0)*1000, 1),
            'progress_values': sorted(set(progress)), 'export_bytes': target.stat().st_size,
            'verified_excel_users': users, 'verified_excel_evidence_rows': raw_rows}


def run(*, users=100, days=365, ui_export=False):
    previous = database.get_db_path()
    with TemporaryDirectory(prefix='komi-shared-report-benchmark-') as root:
        try:
            path = Path(root) / 'synthetic.sqlite3'
            database.set_db_path(path)
            database.init_db()
            database.set_cfg('billing_day', '31')
            with database._open() as conn:
                conn.execute("INSERT INTO devices(slot,name,host,serial,enabled,updated_at) VALUES(0,'Synthetic','192.0.2.10','SYNTHETIC-ONLY',1,'2025-01-01')")
            origin = datetime(2025, 1, 1, 10, tzinfo=UTC)
            def reading(day):
                counters = [WebCounterValue('counter.user.total.total', 'Total', 1000 + day * (n + 1),
                    'User', f'user[{n}].total', Confidence.VERIFIED, 'user', str(n), f'Synthetic user {n}') for n in range(users)]
                counters.append(WebCounterValue('counter.total.total', 'Total', 100000 + day * users,
                    'Device', 'device.total', Confidence.VERIFIED))
                at = origin + timedelta(days=day)
                return at, MeterReading(host='192.0.2.10', collected_at=at.isoformat(), counters=tuple(counters),
                                       source='synthetic-benchmark', auth_state=WebAuthState.ADMIN)
            store = WebObservationStore(path)
            at, r = reading(0)
            anchor = store.record(0, r, serial='SYNTHETIC-ONLY', printer_name='Synthetic', at=at)
            epochs = SlotEpochStore(path)
            epoch = epochs.confirm(epochs.preview(0, anchor['id'], operation='initial', actor='benchmark', reason='Synthetic baseline'))
            began = perf_counter()
            for day in range(1, days):
                at, r = reading(day)
                store.record(0, r, serial='SYNTHETIC-ONLY', printer_name='Synthetic', at=at, epoch_id=epoch)
            indexing = perf_counter() - began
            started = perf_counter()
            with database.report_read_database(path, cancelled=Event()):
                subjects = load_subject_history([0], only_periods=((date(2025, 1, 1), date(2025, 1, 31)),
                                                                  (date(2025, 12, 1), date(2025, 12, 31))))
            subjects_seconds = perf_counter() - started
            started = perf_counter()
            with database.report_read_database(path, cancelled=Event()):
                raw = load_raw_counter_history([0], period_start=date(2025, 1, 1), period_end=date(2025, 12, 31))
            raw_seconds = perf_counter() - started
            started = perf_counter()
            snapshot = load_report_snapshot(0, ('2025-12-01', '2025-12-31'), cache=None, subject_cache=None,
                repository=DeviceConfigRepository(), cancelled=Event(), database_path=path)
            snapshot_seconds = perf_counter() - started
            result = {'users': users, 'days': days, 'index_seconds': round(indexing, 3),
                    'subject_two_periods_seconds': round(subjects_seconds, 3), 'subject_rows': len(subjects),
                    'raw_year_seconds': round(raw_seconds, 3), 'raw_rows': len(raw),
                    'overview_snapshot_seconds': round(snapshot_seconds, 3),
                    'overview_users': len(snapshot.data['subjects']), 'database_bytes': path.stat().st_size}
            if ui_export:
                result.update(measure_export(path, root, users=users, days=days, raw_rows=len(raw)))
            return result
        finally:
            database.set_db_path(previous)


if __name__ == '__main__':
    import argparse
    import os
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ui-export', action='store_true', help='Also measure Qt responsiveness and verify a real year Excel export')
    args = parser.parse_args()
    if args.ui_export:
        os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    print(json.dumps(run(ui_export=args.ui_export), indent=2))
